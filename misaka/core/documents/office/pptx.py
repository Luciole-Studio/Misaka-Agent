"""A deck read shape by shape: text with deviation-only formatting, tables, chart values,
arrow flow, and a mark on what only a vision model could read.

Ported from FrontierAgent's ``plugins/tools/_reader_pptx.py`` (audit D46-D71), the only
reader either project has that treats a slide as a diagram rather than a bag of strings.
What it does that a text dump does not: a chart's real numbers (they live in chart1.xml and
no text extraction reaches them), a colour block and the label drawn on top of it merged
into one line, connectors resolved into ``A -> B -> C``, repeated styling hoisted out of a
row of identical boxes, and an explicit mark on the slides whose meaning is only in the
picture.

Conventions, borrowed whole: a backtick span is parser-added and not file content, and only
deviations are recorded -- a run at the shape's own font size gets no size note.

Three deliberate differences from FrontierAgent:

* the slide marker is ``## Slide N: title``, not ``<!-- slide N -->``. misaka pins a
  citation to a page and ``office.paging`` cuts a deck at that heading, so the marker is
  load-bearing: it is both the boundary and the name the page number resolves to;
* **nothing the parser adds is inserted inside a line.** Emphasis, a run's font-size or
  colour deviation, a table's column separator: all of it goes after the line, or between
  cells as a tab rather than a ``|``. A quotation is checked against the page with every
  whitespace character folded away (``index.normalize_for_quote_match``), so a ``**`` or a
  backtick note wedged between two words means the sentence around it can never be locked
  to its slide -- and ``|`` survives the fold while the spaces around it do not. FA writes
  ``**Total**`` and ``| a | b |`` because that is what markdown looks like; a corpus that
  exists to verify quotations cannot spend them on typography. Emphasis covering a whole
  line is kept, since markers at the two ends cost a within-line quotation nothing, and
  emphasis on a span inside a line moves to a trailing note naming the words that carried
  it. This is also what misaka's HTML and EPUB extractor already does;
* FA closes the document with a list of the slides needing a vision model. Here each slide
  carries its own ``needs_vlm: true``, because a trailing list would be paged into the last
  slide's block and read as a statement about that slide.
"""
from __future__ import annotations

import os

# Text shapes sharing this many backtick bits collapse into one group, with the shared bits
# hoisted to the group header. Below three, unrelated shapes that merely happen to share a
# font size would be merged.
GROUP_MIN_COMMON = 3

# The tolerance, in inches, for "this box contains that label" and "this arrow ends at that
# shape's edge". A deck is drawn by hand, so nothing lines up exactly.
CONTAIN_TOLERANCE = 0.2
EDGE_TOLERANCE = 0.35
LABEL_CHARS = 24

# Beyond this many characters of echoed run text on one line, the trailing note names the
# deviations without repeating the words.
MAX_RUN_ECHO = 200

_PLAIN = (False, False, False)
_MARK_NAMES = ("bold", "italic", "underline")

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P_CNVPR = "{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr"
_A_CNVPR = "{http://schemas.openxmlformats.org/drawingml/2006/main}cNvPr"

_EMU_PER_INCH = 914400

# What python-pptx raises when an accessor does not apply to the shape at hand: a text box
# has no ``.table``, a connector has no ``.auto_shape_type``, a theme-inherited colour has
# no ``.rgb``. Probing is the only way to ask. FrontierAgent catches bare ``Exception`` at
# every one of these; naming the range instead keeps a real bug in this module from being
# swallowed along with them.
_ABSENT = (AttributeError, IndexError, KeyError, TypeError, ValueError)


def _inches(value):
    return round(value / _EMU_PER_INCH, 2) if value is not None else None


def _num(value):
    """28.0 -> '28', 13.5 -> '13.5'. A trailing '.0' on every font size is noise."""
    return f"{value:g}"


def _tag(bits):
    """Meta bits as one backtick span, or "" when there is nothing to say."""
    return "`" + " ".join(bits) + "`" if bits else ""


def _clean(text):
    """Text safe to put in a tab-separated cell."""
    return text.replace("\t", " ").replace("\n", " ").strip()


def _run_color(run):
    try:
        color = run.font.color
        if color is not None and color.type is not None:
            return str(color.rgb)
    except _ABSENT:
        pass
    return ""


def _run_spacing(run):
    """Character spacing from ``a:rPr@spc``, in points."""
    try:
        properties = run._r.find(f"{_A}rPr")
        if properties is not None and properties.get("spc"):
            return round(int(properties.get("spc")) / 100, 1)
    except _ABSENT:
        pass
    return None


def _run_marks(run):
    font = run.font
    return (bool(font.bold), bool(font.italic), bool(font.underline))


def _wrap(text, marks):
    """Apply marks to a whole line, keeping the padding outside the markers."""
    core = text.strip()
    if not core:
        return text
    lead, trail = text[: len(text) - len(text.lstrip())], text[len(text.rstrip()):]
    bold, italic, underline = marks
    if bold and italic:
        core = f"***{core}***"
    elif bold:
        core = f"**{core}**"
    elif italic:
        core = f"*{core}*"
    if underline:
        core = f"<u>{core}</u>"
    return lead + core + trail


def _run_deviations(run, base_size, base_color):
    """What this run does that the shape's default does not: size, colour, spacing.

    Only deviations. The shape's own size and colour are stated once in its meta, so a
    slide of ordinary body text carries no per-run noise at all.
    """
    notes = []
    try:
        size = round(run.font.size.pt, 1) if run.font.size is not None else None
    except _ABSENT:
        size = None
    if size is not None and size != base_size:
        notes.append(f"{_num(size)}pt")
    color = _run_color(run)
    if color and color != "000000" and color != base_color:
        notes.append(f"#{color}")
    spacing = _run_spacing(run)
    if spacing:
        notes.append(f"spc{_num(spacing)}")
    return notes


def _run_note(runs):
    """The trailing note for a line whose runs are not uniform.

    Names each deviation and echoes the words that carried it, so "which phrase was bold"
    and "which one was 14pt" survive -- they simply move off the sentence, where they would
    have broken the quotation.
    """
    by_label = {}
    for text, marks, notes in runs:
        if not text.strip():
            continue
        labels = [name for name, on in zip(_MARK_NAMES, marks, strict=True) if on] + notes
        for label in labels:
            by_label.setdefault(label, []).append(text.strip())
    if not by_label:
        return ""
    echoed = sum(len(word) for words in by_label.values() for word in words)
    if echoed > MAX_RUN_ECHO:
        return "`" + ", ".join(by_label) + "`"
    return "`" + "; ".join(f'{label} "{" / ".join(words)}"'
                           for label, words in by_label.items()) + "`"


def _base_format(frame):
    """The shape's main font size and colour: the most frequent value across its runs.

    A black majority is not hoisted -- black is the default and annotating it would put a
    colour note on every ordinary slide. Only a real non-black majority (an all-white
    title, say) is worth stating once at shape level.
    """
    from collections import Counter
    sizes, colors = [], []
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            try:
                if run.font.size is not None:
                    sizes.append(round(run.font.size.pt, 1))
            except _ABSENT:
                pass
            colors.append(_run_color(run) or "000000")
    base_size = Counter(sizes).most_common(1)[0][0] if sizes else None
    base_color = Counter(colors).most_common(1)[0][0] if colors else None
    return base_size, (None if base_color == "000000" else base_color)


def _list_kind(paragraph):
    """The paragraph's bullet: ``("auto", type, startAt)`` / ``("char", symbol)`` /
    ``("none",)`` / ``None`` when inherited.

    Read from ``a:buAutoNum`` / ``a:buChar`` / ``a:buNone`` rather than guessed from the
    placeholder -- guessing loses level-0 bullets and every numbered list there is.
    """
    try:
        properties = paragraph._p.find(f"{_A}pPr")
    except _ABSENT:
        return None
    if properties is None:
        return None
    if properties.find(f"{_A}buNone") is not None:
        return ("none",)
    auto = properties.find(f"{_A}buAutoNum")
    if auto is not None:
        return ("auto", auto.get("type", "arabicPeriod"), auto.get("startAt"))
    char = properties.find(f"{_A}buChar")
    if char is not None:
        return ("char", char.get("char", "•"))
    return None


def _placeholder(shape):
    if not shape.is_placeholder:
        return ""
    try:
        return str(shape.placeholder_format.type).split()[0]
    except _ABSENT:
        return ""


def _fill(shape):
    """Solid fill colour as RRGGBB, or ""."""
    try:
        fill = shape.fill
        if fill.type is not None and int(fill.type) == 1:  # MSO_FILL.SOLID
            return str(fill.fore_color.rgb)
    except _ABSENT:
        pass
    return ""


def _vertical(shape):
    """The text direction from ``a:bodyPr@vert``; "" when horizontal.

    The text itself reads the same either way; this only says the slide draws it rotated.
    """
    try:
        body = shape.text_frame._txBody.find(f"{_A}bodyPr")
        if body is not None:
            direction = body.get("vert")
            if direction and direction != "horz":
                return direction
    except _ABSENT:
        pass
    return ""


def _bits(shape, with_size=False):
    """Shape-level meta, carrying only what has signal.

    Dropped as zero-information: the AUTO_SHAPE type name, RECTANGLE geometry, Word's
    auto-generated shape names ("Text 3"), and the dimensions of a text box. Kept:
    position, because left-versus-right is how a diagram means anything; the placeholder
    type; geometry when it is not a rectangle; a non-black fill.
    """
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    bits = []
    try:
        if shape.left is not None:
            bits.append(f"@{_inches(shape.left)},{_inches(shape.top)}")
    except _ABSENT:
        pass
    if with_size:
        try:
            if shape.width is not None:
                bits.append(f"{_inches(shape.width)}×{_inches(shape.height)}in")
        except _ABSENT:
            pass
    placeholder = _placeholder(shape)
    if placeholder:
        bits.append(f"ph={placeholder}")
    try:
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and shape.auto_shape_type is not None:
            geometry = str(shape.auto_shape_type).split()[0]
            if geometry != "RECTANGLE":
                bits.append(geometry)
    except _ABSENT:
        pass
    fill = _fill(shape)
    if fill and fill != "000000":
        bits.append(f"fill:#{fill}")
    return bits


def _text_md(shape):
    """A text shape as ``(shape-level format bits, lines)``.

    Bullets and numbering come from the paragraph's own properties; an inherited paragraph
    falls back to the placeholder default -- a body placeholder bullets, a free text box at
    level 0 does not.
    """
    frame = shape.text_frame
    base_size, base_color = _base_format(frame)
    fmt = []
    if base_size is not None:
        fmt.append(f"{_num(base_size)}pt")
    if base_color:
        fmt.append(f"#{base_color}")
    vertical = _vertical(shape)
    if vertical:
        fmt.append(vertical)
    placeholder = _placeholder(shape)
    is_body = any(kind in placeholder for kind in ("BODY", "OBJECT", "SUBTITLE"))
    lines, counters = [], {}
    for paragraph in frame.paragraphs:
        runs = [(run.text, _run_marks(run), _run_deviations(run, base_size, base_color))
                for run in paragraph.runs if run.text]
        text = "".join(part for part, _, _ in runs).strip() or paragraph.text.strip()
        if not text:
            continue
        # Uniform across the line: the markers land at the two ends, where a quotation of
        # anything inside the line steps over neither. Otherwise the whole thing moves to a
        # trailing note -- see the module docstring.
        styles = {(marks, tuple(notes)) for part, marks, notes in runs if part.strip()}
        if len(styles) == 1:
            marks, notes = styles.pop()
            if any(marks):
                text = _wrap(text, marks)
            if notes:
                text = f"{text} `{','.join(notes)}`"
        elif (note := _run_note(runs)):
            text = f"{text} {note}"
        level = paragraph.level
        indent = "  " * level
        kind = _list_kind(paragraph)
        if kind and kind[0] == "auto":
            start = int(kind[2]) if len(kind) > 2 and kind[2] else 1
            ordinal = counters.get(level, start)
            counters[level] = ordinal + 1
            lines.append(f"{indent}{ordinal}. {text}")
        elif kind and kind[0] == "char":
            lines.append(f"{indent}- {text}")
        elif kind and kind[0] == "none":
            lines.append(f"{indent}{text}")
        elif is_body or level > 0:
            lines.append(f"{indent}- {text}")
        else:
            lines.append(text)
    return fmt, lines


def _table_md(shape):
    """A table as ``(rows, columns, tab-separated lines)``.

    A merged cell repeats the origin's text in the file; only the origin keeps it, so one
    value does not read as several.
    """
    table = shape.table
    rows, columns = len(table.rows), len(table.columns)
    lines = []
    for r in range(rows):
        cells = []
        for c in range(columns):
            cell = table.cell(r, c)
            cells.append("" if cell.is_spanned else _clean(cell.text))
        lines.append("\t".join(cells))
    return rows, columns, lines


def _custom_labels(chart):
    """Data-label text that overrides the plotted number.

    An automatic label just repeats the series value, which is already in the data rows, so
    only an overridden one is worth carrying. Every plot is scanned: a combo chart has more
    than one.
    """
    found = {}
    try:
        for plot in chart.plots:
            for series in plot.series:
                for index, point in enumerate(series.points):
                    try:
                        label = point.data_label
                        if label.has_text_frame and label.text_frame.text.strip():
                            found.setdefault(str(series.name), {})[index] = \
                                label.text_frame.text.strip()
                    except _ABSENT:
                        pass
    except _ABSENT:
        pass
    return found or None


def _chart_md(shape):
    """A chart as ``(meta bits, lines)``, with its real values.

    The numbers live in the embedded chart part and no text extraction reaches them, so
    this goes through the chart API: categories down, series across.
    """
    chart = shape.chart
    bits = ["chart", str(chart.chart_type).split()[0] if chart.chart_type is not None else "?"]
    if chart.has_title:
        title_text = " ".join(chart.chart_title.text_frame.text.split()).replace('"', "'")
        bits.append(f'title="{title_text}"')
    try:
        if chart.plots[0].has_data_labels:
            bits.append("data-labels-shown")
    except _ABSENT:
        pass
    lines = []
    try:
        series, categories = [], []
        for plot in chart.plots:
            try:
                plot_categories = list(plot.categories)
                if plot_categories and not categories:
                    categories = plot_categories
            except _ABSENT:
                pass
            for one in plot.series:
                series.append((str(one.name), list(one.values)))
        if categories:
            lines.append("\t".join(["category"] + [name for name, _ in series]))
            for index, category in enumerate(categories):
                row = [_clean(str(category))]
                row += [str(values[index]) if index < len(values) else ""
                        for _, values in series]
                lines.append("\t".join(row))
        elif series:  # scatter and bubble have no category axis
            for name, values in series:
                lines.append(f'- series "{name}": ' + ", ".join(str(v) for v in values))
        custom = _custom_labels(chart)
        if custom:
            lines.append(f"`custom data-label text: {custom}`")
    except _ABSENT as error:
        lines.append(f"`chart data unavailable: {type(error).__name__}: {error}`")
    return bits, lines


def _picture_alt(shape):
    """A picture's alt text -- the only textual description it has."""
    try:
        # An lxml element with no children is falsy, so ``a or b`` would silently discard a
        # childless cNvPr; the check has to be ``is None``.
        node = shape._element.find(f".//{_P_CNVPR}")
        if node is None:
            node = shape._element.find(f".//{_A_CNVPR}")
        if node is not None:
            return node.get("descr") or ""
    except _ABSENT:
        pass
    return ""


def _bbox(shape):
    """``(left, top, width, height)`` in inches, or None."""
    try:
        if shape.left is None:
            return None
        return (_inches(shape.left), _inches(shape.top),
                _inches(shape.width), _inches(shape.height))
    except _ABSENT:
        return None


def _shape_text(shape):
    try:
        return shape.text_frame.text.strip() if getattr(shape, "has_text_frame", False) else ""
    except _ABSENT:
        return ""


def _is_line(shape):
    """Whether this shape is a connector: what matters is the connection, not the drawing.

    A block arrow *with* text is not one -- it is a labelled step, and its label is content.
    """
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    try:
        if shape.shape_type == MSO_SHAPE_TYPE.LINE:
            return True
    except _ABSENT:
        pass
    box = _bbox(shape)
    if box and (box[2] <= 0.05 or box[3] <= 0.05):
        return True
    try:
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            try:
                kind = str(shape.auto_shape_type).upper()
            except _ABSENT as error:
                # A 'line' preset raises on auto_shape_type; the message names it.
                return "line" in str(error).lower()
            if "ARROW" in kind and not _shape_text(shape):
                return True
    except _ABSENT:
        pass
    return False


def _is_box(shape):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    try:
        return shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and not _is_line(shape)
    except _ABSENT:
        return False


def _pair_label_boxes(shapes, boxes):
    """Match each text shape to the empty box drawn behind it.

    Authors draw a colour block and its caption as two overlapping shapes; emitting both
    puts an empty backtick span above every label on the slide. The container must enclose
    the label on all four sides and not be absurdly larger than it, and the tightest fit
    wins. Returns ``({text index: box index}, {box indices consumed})``.
    """
    containers = [i for i, shape in enumerate(shapes)
                  if _is_box(shape) and not _shape_text(shape) and boxes[i]]
    pairs, used = {}, set()
    for i, shape in enumerate(shapes):
        if not _shape_text(shape) or not boxes[i]:
            continue
        left, top, width, height = boxes[i]
        best, best_area = None, None
        for index in containers:
            if index in used:
                continue
            cl, ct, cw, ch = boxes[index]
            if (cl - CONTAIN_TOLERANCE <= left and ct - CONTAIN_TOLERANCE <= top
                    and cl + cw + CONTAIN_TOLERANCE >= left + width
                    and ct + ch + CONTAIN_TOLERANCE >= top + height
                    and cw * ch <= 10 * max(width * height, 0.01)):
                area = cw * ch
                if best is None or area < best_area:
                    best, best_area = index, area
        if best is not None:
            pairs[i] = best
            used.add(best)
    return pairs, used


def _label_at_edge(point, edge, shapes, boxes):
    """The text of the shape whose named edge sits against ``point``.

    This is how an arrow finds what it joins: at the arrow's left end, look for a shape
    whose *right* edge is there.
    """
    px, py = point
    for i, shape in enumerate(shapes):
        box, text = boxes[i], _shape_text(shape)
        if not box or not text:
            continue
        x0, top, width, height = box
        if edge in ("right", "left") and not (top - 0.2 <= py <= top + height + 0.2):
            continue
        if edge in ("top", "bottom") and not (x0 - 0.2 <= px <= x0 + width + 0.2):
            continue
        distance = {"right": abs((x0 + width) - px), "left": abs(x0 - px),
                    "bottom": abs((top + height) - py), "top": abs(top - py)}[edge]
        if distance <= EDGE_TOLERANCE:
            return text[:LABEL_CHARS]
    return None


def _arrow_info(shape, index, shapes, boxes):
    """A connector resolved into direction and the two labels it joins."""
    box = boxes[index]
    if not box:
        return {"line": _tag(["line", "vlm"]), "src": None, "dst": None,
                "arrow": "→", "connected": False}
    x0, top, width, height = box
    try:
        kind = str(shape.auto_shape_type).upper()
    except _ABSENT:
        kind = ""
    left_end, right_end = (x0, top + height / 2), (x0 + width, top + height / 2)
    top_end, bottom_end = (x0 + width / 2, top), (x0 + width / 2, top + height)
    if width >= height and not ("UP_ARROW" in kind or "DOWN_ARROW" in kind):
        if "LEFT_ARROW" in kind and "RIGHT" not in kind:
            src = _label_at_edge(right_end, "left", shapes, boxes)
            dst = _label_at_edge(left_end, "right", shapes, boxes)
            arrow = "←"
        else:  # a plain connector has no direction of its own; left to right is the default
            src = _label_at_edge(left_end, "right", shapes, boxes)
            dst = _label_at_edge(right_end, "left", shapes, boxes)
            arrow = "→"
    elif "UP_ARROW" in kind:
        src = _label_at_edge(bottom_end, "top", shapes, boxes)
        dst = _label_at_edge(top_end, "bottom", shapes, boxes)
        arrow = "↑"
    else:
        src = _label_at_edge(top_end, "bottom", shapes, boxes)
        dst = _label_at_edge(bottom_end, "top", shapes, boxes)
        arrow = "↓"
    bits = [f"@{x0},{top}", f"arrow {arrow}"]
    if src or dst:
        bits.append(f'connects "{src or "?"}"{arrow}"{dst or "?"}"')
    else:
        bits.append("vlm")
    return {"line": _tag(bits), "src": src, "dst": dst, "arrow": arrow,
            "connected": bool(src and dst)}


def _order_flow(edges):
    """Chain edges into ``A -> B -> C``; anything that is not one path lists its edges."""
    if not edges:
        return None
    successor = dict(edges)
    targets = {dst for _, dst in edges}
    starts = [src for src, _ in edges if src not in targets]
    nodes = {src for src, _ in edges} | targets
    if len(starts) == 1:
        chain, seen = [starts[0]], {starts[0]}
        while chain[-1] in successor and successor[chain[-1]] not in seen:
            chain.append(successor[chain[-1]])
            seen.add(chain[-1])
        if len(chain) == len(nodes):
            return " → ".join(chain)
    return ", ".join(f"{src}→{dst}" for src, dst in edges)


def _records(shapes, boxes, pairs, used, skip):
    """One record per shape: what it is and the lines it renders to."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE as MSO
    records = []
    for index, shape in enumerate(shapes):
        if index in used or index in skip:
            continue
        record = {"idx": index, "kind": "block", "lines": []}
        try:
            if getattr(shape, "has_table", False):
                rows, columns, table = _table_md(shape)
                record["lines"] = [_tag([*_bits(shape), f"table {rows}×{columns}"]), *table]
            elif getattr(shape, "has_chart", False):
                chart_bits, chart_lines = _chart_md(shape)
                record["lines"] = [_tag(_bits(shape) + chart_bits), *chart_lines]
            elif shape.shape_type == MSO.PICTURE:
                alt = _picture_alt(shape)
                bits = [*_bits(shape, with_size=True), "image", "vlm"]
                bits.append(f'alt:"{alt}"' if alt else "no-alt")
                record["lines"] = [_tag(bits)]
                record["needs_vlm"] = True
            elif shape.shape_type == MSO.GROUP:
                # A group is left whole: decomposing it loses the alignment that made the
                # author group it, and the slide goes to a vision model anyway.
                try:
                    count = len(shape.shapes)
                except _ABSENT:
                    count = "?"
                record["lines"] = [_tag([*_bits(shape, with_size=True),
                                         f"group-of-{count}-shapes", "not-decomposed", "vlm"])]
                record["needs_vlm"] = True
            elif _is_line(shape):
                record["kind"] = "arrow"
                record["info"] = _arrow_info(shape, index, shapes, boxes)
            elif shape.has_text_frame and shape.text_frame.text.strip():
                fmt, lines = _text_md(shape)
                box = pairs.get(index)
                meta = (_bits(shapes[box], with_size=True) if box is not None
                        else _bits(shape))
                position = meta[0] if meta and meta[0].startswith("@") else ""
                rest = [bit for bit in meta if not bit.startswith("@")]
                record.update(kind="text", meta=meta, fmt=fmt, lines=lines, pos=position,
                              sig=tuple(rest + fmt), label=_shape_text(shape))
            else:
                record["lines"] = [_tag([*_bits(shape, with_size=True), "vlm"])]
                record["needs_vlm"] = True
        except Exception as error:  # noqa: BLE001 - one bad shape must not lose the slide
            record["lines"] = [f"`shape parse error: {type(error).__name__}: {error}`"]
        records.append(record)
    return records


def _groups(records):
    """Cluster text shapes that share styling into groups.

    A deck's five feature boxes are one object drawn five times; printing five identical
    backtick spans buries the five labels that actually differ. Two shapes join when they
    share at least ``GROUP_MIN_COMMON`` bits, and a cluster survives only if the
    intersection across *all* its members is still that large -- otherwise a chain of
    pairwise matches produces a group with nothing in common.
    """
    from collections import defaultdict
    text_records = [r for r in records if r["kind"] == "text"]
    for record in text_records:
        record["nonpos"] = [b for b in record["meta"] if not b.startswith("@")] + record["fmt"]
        record["bitset"] = set(record["nonpos"])
    parent = {r["idx"]: r["idx"] for r in text_records}

    def find(node):
        while parent[node] != node:
            parent[node] = parent[parent[node]]
            node = parent[node]
        return node

    for a in range(len(text_records)):
        for b in range(a + 1, len(text_records)):
            if len(text_records[a]["bitset"] & text_records[b]["bitset"]) >= GROUP_MIN_COMMON:
                root_a, root_b = find(text_records[a]["idx"]), find(text_records[b]["idx"])
                if root_a != root_b:
                    parent[root_a] = root_b
    clusters = defaultdict(list)
    for record in text_records:
        clusters[find(record["idx"])].append(record)
    members_by_root, root_by_record = {}, {}
    for root, members in clusters.items():
        if len(members) < 2:
            continue
        if len(set.intersection(*[m["bitset"] for m in members])) < GROUP_MIN_COMMON:
            continue
        members_by_root[root] = members
        for member in members:
            root_by_record[member["idx"]] = root
    return members_by_root, root_by_record


def _slide_lines(shapes, boxes, pairs, used, skip):
    """One slide's body lines, and whether it needs a vision model."""
    from collections import defaultdict
    records = _records(shapes, boxes, pairs, used, skip)
    members_by_root, root_by_record = _groups(records)

    label_to_group = {}
    for root, members in members_by_root.items():
        for member in members:
            if member["label"]:
                label_to_group[member["label"][:LABEL_CHARS]] = root

    group_edges = defaultdict(list)
    for record in records:
        if record["kind"] != "arrow":
            continue
        src, dst = record["info"]["src"], record["info"]["dst"]
        if src and dst and label_to_group.get(src) and label_to_group.get(src) == label_to_group.get(dst):
            group_edges[label_to_group[src]].append((src, dst))
            record["consumed"] = True
        elif not record["info"]["connected"]:
            record["needs_vlm"] = True

    out, emitted, visual = [], set(), []
    for record in records:
        if record.get("consumed"):
            continue
        if record.get("needs_vlm"):
            # Purely visual elements are gathered at the end rather than interleaved: they
            # carry no text, so in reading order they only interrupt what does.
            visual.extend([record["info"]["line"]] if record["kind"] == "arrow"
                          else record["lines"])
            continue
        out.append("")
        if record["kind"] == "text" and record["idx"] in root_by_record:
            root = root_by_record[record["idx"]]
            if root in emitted:
                out.pop()
                continue
            emitted.add(root)
            members = members_by_root[root]
            common = [bit for bit in members[0]["nonpos"]
                      if all(bit in other["bitset"] for other in members)]
            out.append(_tag([f"group ×{len(members)}", *common]))
            for member in members:
                extra = [member["pos"]] + [b for b in member["nonpos"] if b not in common]
                member_tag = _tag([bit for bit in extra if bit])
                if len(member["lines"]) == 1:
                    out.append(f"- {member_tag} {member['lines'][0]}".strip())
                else:
                    out.append(f"- {member_tag}".rstrip())
                    out.extend("  " + line for line in member["lines"])
            flow = _order_flow(group_edges.get(root, []))
            if flow:
                out.append(f"- flow: {flow}")
        elif record["kind"] == "text":
            tag = _tag(record["meta"] + record["fmt"])
            if len(record["lines"]) == 1 and tag:
                out.append(f"{tag} {record['lines'][0]}")
            else:
                if tag:
                    out.append(tag)
                out.extend(record["lines"])
        elif record["kind"] == "arrow":
            out.append(record["info"]["line"])
        else:
            out.extend(record["lines"])
    if visual:
        out += ["", "▸ visual elements — needs_vlm: true", *visual]
    return out, bool(visual)


def _slide_title(shapes):
    """``(title text, shape index)`` for the slide's title placeholder, else ``("", None)``.

    The title goes into the ``## Slide N:`` line, which is both the page boundary and the
    name a citation's page number resolves to, so the shape is not rendered again below it.
    """
    for index, shape in enumerate(shapes):
        if "TITLE" not in _placeholder(shape):
            continue
        text = _shape_text(shape)
        if text:
            return " ".join(text.split()), index
    return "", None


def title(path):
    """The deck's own title: the first slide's title placeholder, or ""."""
    try:
        from pptx import Presentation
        presentation = Presentation(str(path))
        for slide in presentation.slides:
            text, _ = _slide_title(list(slide.shapes))
            return text
    except Exception:  # noqa: BLE001 - a title is a nicety; the file name is the fallback
        return ""
    return ""


def render(path):
    """The whole deck as markdown, one ``## Slide N`` block per slide.

    Raises ``ValueError`` naming the file when python-pptx cannot open it -- that string is
    what ``doc_add`` hands the model.
    """
    try:
        from pptx import Presentation
    except ImportError as error:                                    # pragma: no cover
        raise ValueError(
            f"Cannot read {os.path.basename(str(path))}: python-pptx is not installed."
        ) from error
    try:
        presentation = Presentation(str(path))
    except Exception as error:
        raise ValueError(
            f"Cannot read {os.path.basename(str(path))}: {type(error).__name__}: {error}"
        ) from error

    out = [
        (f"<!-- pptx readout · slide "
         f"{_inches(presentation.slide_width)}×{_inches(presentation.slide_height)}in. "
         "A backtick span is parser-added, not slide content: @x,y is position in inches, "
         "NNpt/#hex are font size and colour and appear only where they deviate from the "
         "shape default, arrow/connects is a connector, vlm marks a purely visual element. "
         "Table and chart rows are tab-separated. -->"),
    ]
    for number, slide in enumerate(presentation.slides, 1):
        shapes = list(slide.shapes)
        heading, title_index = _slide_title(shapes)
        out.append("")
        out.append(f"## Slide {number}: {heading}" if heading else f"## Slide {number}")
        boxes = [_bbox(shape) for shape in shapes]
        pairs, used = _pair_label_boxes(shapes, boxes)
        skip = {title_index} if title_index is not None else set()
        lines, _needs_vlm = _slide_lines(shapes, boxes, pairs, used, skip)
        out.extend(lines)
        if slide.has_notes_slide:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                # At the end of the slide, not the top: the speaker's aside should not sit
                # between the slide's title and the slide's own words.
                out.append("")
                out.append("> **notes:** " + note.replace("\n", "\n> "))
    while out and not out[-1].strip():
        out.pop()
    return "\n".join(out) + "\n"
