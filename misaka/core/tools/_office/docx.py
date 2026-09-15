"""Writing a Word document: thirteen ops, applied in order to one file.

Ported from FrontierAgent's ``plugins/tools/_writer_docx.py`` (audit D85-D103). The
addressing is the part worth keeping: an op names its target by **content** -- ``find`` and
``after_text`` -- not by a positional index. A model that has to count paragraphs is a model
that edits the wrong one as soon as an earlier op inserts anything.

One deliberate improvement over FrontierAgent. Its ``_replace_in_paragraph`` rewrites the
whole paragraph into its first run and blanks the rest, so a single find-and-replace strips
the bold off every other word in that paragraph -- silently, with a success receipt. Here
the replacement is spliced by character offset: the runs outside the match are not touched
at all, and the match itself keeps the formatting of the run it started in.
"""
from __future__ import annotations

import os

from misaka.core.tools._office._receipt import result
from misaka.core.tools._office._runs import norm_runs, splice_runs

SUFFIXES = frozenset({".docx"})

_ALIGN = {"left": "LEFT", "center": "CENTER", "centre": "CENTER",
          "right": "RIGHT", "justify": "JUSTIFY"}

# Word's page-number field formats, by the names the tool contract uses.
_PAGE_FORMATS = {"decimal": "decimal", "roman_lower": "lowerRoman",
                 "roman_upper": "upperRoman", "alpha": "lowerLetter"}

# The colour Word itself uses for a hyperlink.
_LINK_COLOR = "0563C1"

OPS = ("create", "replace_text", "insert_paragraph", "insert_heading", "insert_table",
       "format_text", "format_paragraph", "add_hyperlink", "add_image",
       "set_page_number", "set_page_margins", "set_page_orientation", "set_header_footer")


def _align(paragraph, value):
    if not value:
        return
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    name = _ALIGN.get(str(value).lower())
    if name:
        paragraph.alignment = getattr(WD_ALIGN_PARAGRAPH, name)


def _hyperlink(paragraph, url, text, *, bold=None, italic=None, color=None, underline=True):
    """Append a real hyperlink run. python-docx has no API for one, so this is OXML."""
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    relation = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), str(color or _LINK_COLOR).lstrip("#"))
    properties.append(colour)
    if underline:
        underline_node = OxmlElement("w:u")
        underline_node.set(qn("w:val"), "single")
        properties.append(underline_node)
    if bold:
        properties.append(OxmlElement("w:b"))
    if italic:
        properties.append(OxmlElement("w:i"))
    run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = text or url
    text_node.set(qn("xml:space"), "preserve")
    run.append(text_node)
    link.append(run)
    paragraph._p.append(link)
    return link


def _apply_runs(paragraph, text):
    """Write rich text into a paragraph, one run per formatting change."""
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    for spec in norm_runs(text):
        if spec.get("link"):
            _hyperlink(paragraph, spec["link"], spec.get("text", ""),
                       bold=spec.get("bold"), italic=spec.get("italic"),
                       color=spec.get("color"), underline=spec.get("underline", True))
            continue
        run = paragraph.add_run(spec.get("text", ""))
        for field, attribute in (("bold", "bold"), ("italic", "italic"),
                                 ("underline", "underline")):
            if spec.get(field) is not None:
                setattr(run, attribute, bool(spec[field]))
        if spec.get("strike") is not None:
            run.font.strike = bool(spec["strike"])
        if spec.get("color"):
            run.font.color.rgb = RGBColor.from_string(str(spec["color"]).lstrip("#"))
        if spec.get("size"):
            run.font.size = Pt(float(spec["size"]))
        if spec.get("font"):
            run.font.name = spec["font"]
            # python-docx writes only w:ascii and w:hAnsi. CJK characters take their font
            # from w:eastAsia, so without this ``"font": "SimSun"`` has no effect at all on
            # Chinese, Japanese or Korean text -- the theme font is used and the document
            # comes out in something other than what was asked for.
            run._element.rPr.rFonts.set(qn("w:eastAsia"), spec["font"])


def _list_style(paragraph, spec):
    """Put the paragraph in a list. Never by writing ``- `` into the text."""
    if not spec:
        return
    from docx.shared import Inches
    kind = str(spec.get("type", "bullet")).lower()
    level = int(spec.get("level", 0))
    base = "List Number" if kind.startswith("num") else "List Bullet"
    try:
        paragraph.style = base if level == 0 else f"{base} {min(level + 1, 3)}"
    except KeyError:
        paragraph.style = base
    if level > 0:
        paragraph.paragraph_format.left_indent = Inches(0.25 * (level + 1))


def _flow(paragraph, args):
    """Pagination properties: whether the paragraph starts a page or holds onto the next."""
    fmt = paragraph.paragraph_format
    if args.get("page_break_before") is not None:
        fmt.page_break_before = bool(args["page_break_before"])
    if args.get("keep_with_next") is not None:
        fmt.keep_with_next = bool(args["keep_with_next"])


def _spacing(paragraph, args):
    from docx.shared import Pt
    fmt = paragraph.paragraph_format
    if args.get("line_spacing") is not None:
        fmt.line_spacing = float(args["line_spacing"])
    if args.get("space_before") is not None:
        fmt.space_before = Pt(float(args["space_before"]))
    if args.get("space_after") is not None:
        fmt.space_after = Pt(float(args["space_after"]))


def _fill(cell, color):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), str(color).lstrip("#"))
    properties.append(shading)


def _cell(cell, value):
    """Write one table cell: rich text, or ``{content, bold, align, fill_color}``."""
    content, extra = value, {}
    if isinstance(value, dict) and "content" in value:
        content, extra = value["content"], value
    paragraph = cell.paragraphs[0]
    _apply_runs(paragraph, content)
    if extra.get("bold"):
        for run in paragraph.runs:
            run.bold = True
    if extra.get("align"):
        _align(paragraph, extra["align"])
    if extra.get("fill_color"):
        _fill(cell, extra["fill_color"])


def _header_row(table):
    """Bold the first row and mark it ``w:tblHeader`` so it repeats across a page break."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    properties = table.rows[0]._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def _column_widths(table, widths):
    """Fix the column widths, all four parts of it.

    Word and LibreOffice both default to auto-fit and each reads a different place, so a
    width set only one way is a width that silently does nothing in the other program:
    the table layout must say ``fixed``, ``w:tblGrid`` carries the widths LibreOffice
    reads, every cell's ``tcW`` carries the ones Word reads, and autofit must be off.
    """
    if not widths:
        return
    from docx.oxml.ns import qn
    from docx.shared import Inches
    # The ``autofit`` setter inserts ``w:tblLayout`` where the schema's sequence requires
    # it, before ``w:tblLook``. Appending the element by hand puts it after, which Word
    # reports as a corrupt file (LibreOffice is lenient and would hide the mistake).
    table.autofit = False
    columns = len(table.columns)
    twips = []
    for width in widths[:columns]:
        try:
            twips.append(max(round(float(width) * 1440), 1))
        except (TypeError, ValueError):
            twips.append(None)
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        grid_columns = grid.findall(qn("w:gridCol"))
        for index, value in enumerate(twips):
            if value is not None and index < len(grid_columns):
                grid_columns[index].set(qn("w:w"), str(value))
    # Iterating rows rather than columns: the column iterator repeats merged cells.
    for row in table.rows:
        for index, value in enumerate(twips):
            if value is not None and index < len(row.cells):
                row.cells[index].width = Inches(float(widths[index]))


def _paragraphs(document):
    """Every paragraph in the body and inside table cells.

    Find-and-replace has to reach the tables: a figure quoted in a table is exactly the
    kind of thing that needs correcting, and a walk over ``doc.paragraphs`` never sees it.
    """
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    # XML order visits merged cells once and includes nested tables/content controls.
    for element in document.element.body.iter(qn("w:p")):
        yield Paragraph(element, document)


def _splice(paragraph, find, replace, budget):
    """Replace occurrences of ``find`` in one paragraph, keeping every other run intact.

    Works on the paragraph's character offsets: the replacement goes into the run the match
    starts in, and the matched characters are cut out of the runs that follow. Runs that do
    not overlap the match are never rewritten, so the paragraph's formatting survives a
    correction -- which is the whole difference from rewriting the paragraph into run 0.
    """
    from docx.text.run import Run
    runs = [Run(node, paragraph) for node in paragraph._p.xpath("./w:r|./w:hyperlink/w:r")]
    return splice_runs(runs, find, replace, budget)


def _anchor(document, text):
    for paragraph in document.paragraphs:
        if text in paragraph.text:
            return paragraph
    return None


def _after(anchor, document):
    """A new empty paragraph after ``anchor``, or at the end when there is none."""
    if anchor is None:
        return document.add_paragraph()
    fresh = anchor.insert_paragraph_before()
    anchor._p.addnext(fresh._p)
    return fresh


def _table(document, rows, args):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = args.get("style", "Table Grid")
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            _cell(table.cell(r, c), value)
    if args.get("header", True):
        _header_row(table)
    _column_widths(table, args.get("column_widths_in"))
    if args.get("cant_split"):
        for row in table.rows:
            properties = row._tr.get_or_add_trPr()
            # The schema fixes the order inside ``w:trPr`` and a header row has already
            # appended ``w:tblHeader``; appending here would invert the two and produce a
            # file Word refuses. ``insert_element_before`` also guards a second call.
            if properties.find(qn("w:cantSplit")) is None:
                properties.insert_element_before(
                    OxmlElement("w:cantSplit"),
                    "w:trHeight", "w:tblHeader", "w:tblCellSpacing", "w:jc", "w:hidden")
    return table


def _block(document, spec):
    """Append one content block. Returns its singular label, or ``None`` when it wrote
    nothing -- the caller reports those rather than letting a block vanish."""
    from docx.shared import Inches
    kind = spec.get("type", "paragraph")
    if not spec.get("type") and not any(
            spec.get(field) for field in ("text", "items", "rows", "path", "image_path")):
        # A block with neither a type nor any content is a mistake, not a spacer. Writing
        # an empty paragraph for it and counting it makes ``create`` report success for a
        # document that is blank. An explicit ``{"type": "paragraph"}`` still gets its
        # empty paragraph -- that one was asked for.
        return None
    if kind == "heading":
        heading = document.add_heading("", level=int(spec.get("level", 1)))
        _apply_runs(heading, spec.get("text", ""))
        _align(heading, spec.get("align"))
        return "heading"
    if kind == "paragraph":
        paragraph = document.add_paragraph(style=spec.get("style"))
        _apply_runs(paragraph, spec.get("text", ""))
        _list_style(paragraph, spec.get("list"))
        _align(paragraph, spec.get("align"))
        _flow(paragraph, spec)
        _spacing(paragraph, spec)
        return "paragraph"
    if kind in ("bullet_list", "numbered_list", "list"):
        items = spec.get("items", [])
        if not items:
            return None
        numbered = kind == "numbered_list" or bool(spec.get("ordered") or spec.get("numbered"))
        style = spec.get("style", "List Number" if numbered else "List Bullet")
        for item in items:
            _apply_runs(document.add_paragraph(style=style), item)
        return "list"
    if kind == "table":
        rows = spec.get("rows", [])
        if not rows:
            return None
        _table(document, rows, spec)
        return "table"
    if kind == "image":
        source = spec.get("path") or spec.get("image_path")
        if not source or not os.path.exists(str(source)):
            return None
        options = {}
        if spec.get("width_in"):
            options["width"] = Inches(float(spec["width_in"]))
        if spec.get("height_in"):
            options["height"] = Inches(float(spec["height_in"]))
        document.add_picture(str(source), **options)
        return "image"
    if kind == "page_break":
        document.add_page_break()
        return "page_break"
    # Unknown block type, but it carries content: write it rather than drop it silently.
    if spec.get("items"):
        for item in spec["items"]:
            _apply_runs(document.add_paragraph(style="List Bullet"), item)
        return "list"
    if spec.get("text"):
        _apply_runs(document.add_paragraph(), spec.get("text", ""))
        return "paragraph"
    return None


def _create(path, args, overwrite):
    import docx as _docx
    if os.path.exists(path) and not (overwrite or args.get("overwrite")):
        return result(f"create refused: {path} already exists — use insert_* or replace_text "
                      "to add to it, or pass overwrite:true to rebuild it", ok=False)
    document = _docx.Document()
    properties = document.core_properties
    for field in ("title", "subject", "author", "comments"):
        if (args.get("metadata") or {}).get(field):
            setattr(properties, field, args["metadata"][field])
    blocks = args.get("blocks", [])
    counts, empty = {}, []
    for index, spec in enumerate(blocks):
        label = _block(document, spec)
        if label is None:
            empty.append(index)
        else:
            counts[label] = counts.get(label, 0) + 1
    if not counts:
        return result(f"create wrote nothing to {path}: {len(blocks)} block(s), all empty or "
                      "unknown — a block needs type plus text/items/rows", ok=False)
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    document.save(path)
    return result(f"created docx: {path}", counts=counts,
                  warn=(f"block(s) {empty} wrote nothing (check the schema)" if empty else None))


def _page_number(document, args):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    where = str(args.get("location", "footer")).lower()
    section = document.sections[0]
    container = section.header if where == "header" else section.footer
    container.is_linked_to_previous = False
    paragraph = container.paragraphs[0] if container.paragraphs else container.add_paragraph()
    for run in list(paragraph.runs):
        run.text = ""
    _align(paragraph, args.get("align", "center"))

    def field(instruction):
        run = paragraph.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        run._r.append(begin)
        instruction_node = OxmlElement("w:instrText")
        instruction_node.set(qn("xml:space"), "preserve")
        instruction_node.text = instruction
        run._r.append(instruction_node)
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.append(end)

    if args.get("of_total"):
        field("PAGE")
        paragraph.add_run(" / ")
        field("NUMPAGES")
    else:
        field("PAGE")
    if args.get("start") is not None or args.get("fmt"):
        section_properties = section._sectPr
        numbering = section_properties.find(qn("w:pgNumType"))
        if numbering is None:
            numbering = OxmlElement("w:pgNumType")
            section_properties.append(numbering)
        if args.get("start") is not None:
            numbering.set(qn("w:start"), str(int(args["start"])))
        if args.get("fmt"):
            numbering.set(qn("w:fmt"), _PAGE_FORMATS.get(args["fmt"], "decimal"))
    return f"set page number in {where} (of_total={bool(args.get('of_total'))})"


def write(path, op, args, *, overwrite=False):
    """Apply one op to a .docx. Returns a result dict or an ``[error]`` string."""
    try:
        import docx as _docx
        from docx.shared import Inches, Pt, RGBColor
    except ImportError:                                             # pragma: no cover
        return "[error] python-docx is not installed."

    if op == "create":
        return _create(path, args, overwrite)
    if not os.path.exists(path):
        return f"[error] file not found (this op edits an existing file): {path}"
    document = _docx.Document(path)

    if op == "replace_text":
        find = args.get("find")
        if not find:
            return "[error] replace_text needs 'find'."
        limit = args.get("count")
        budget = [int(limit) if limit is not None else -1]
        if budget[0] == 0:
            return result("replace_text: 0 replacements requested; file unchanged")
        done = 0
        for paragraph in _paragraphs(document):
            if budget[0] == 0:
                break
            done += _splice(paragraph, find, args.get("replace", ""), budget)
        document.save(path)
        return result(f"replaced {done} occurrence(s) of {find!r}",
                      warn=(f"0 matches for {find!r}" if done == 0 else None))

    if op in ("insert_paragraph", "insert_heading", "insert_table", "add_image"):
        after_text = args.get("after_text")
        anchor = _anchor(document, after_text) if after_text else None
        if after_text and anchor is None:
            return result(f"{op} skipped",
                          warn=f"anchor {after_text!r} not found — nothing inserted")
        if op == "insert_paragraph":
            paragraph = _after(anchor, document)
            if args.get("style"):
                paragraph.style = args["style"]
            _apply_runs(paragraph, args.get("text", ""))
            _list_style(paragraph, args.get("list"))
            _align(paragraph, args.get("align"))
            _flow(paragraph, args)
            _spacing(paragraph, args)
            document.save(path)
            return f"inserted paragraph ({'after anchor' if anchor else 'at end'})"
        if op == "insert_heading":
            level = int(args.get("level", 1))
            if anchor is not None:
                paragraph = _after(anchor, document)
                paragraph.style = document.styles[f"Heading {min(level, 9)}"]
                _apply_runs(paragraph, args.get("text", ""))
            else:
                _apply_runs(document.add_heading("", level=level), args.get("text", ""))
            document.save(path)
            return f"inserted heading L{level}"
        if op == "insert_table":
            rows = args.get("rows")
            if not rows:
                return "[error] insert_table needs 'rows'."
            table = _table(document, rows, args)
            if anchor is not None:
                anchor._p.addnext(table._tbl)
            document.save(path)
            return f"inserted table {len(rows)}×{len(rows[0])}"
        source = args.get("image_path") or args.get("path")
        if not source or not os.path.exists(str(source)):
            return f"[error] add_image: no such image {source!r}"
        options = {}
        if args.get("width"):
            options["width"] = Inches(float(args["width"]))
        if args.get("height"):
            options["height"] = Inches(float(args["height"]))
        if anchor is not None:
            _after(anchor, document).add_run().add_picture(str(source), **options)
        else:
            document.add_picture(str(source), **options)
        document.save(path)
        return f"added image {os.path.basename(str(source))}"

    if op == "format_text":
        find = args.get("find")
        if not find:
            return "[error] format_text needs 'find'."
        hits = 0
        for paragraph in _paragraphs(document):
            if find not in paragraph.text:
                continue
            for run in paragraph.runs:
                if not ((run.text and run.text in find) or find in (run.text or "")):
                    continue
                for field, attribute in (("bold", "bold"), ("italic", "italic"),
                                         ("underline", "underline")):
                    if args.get(field) is not None:
                        setattr(run, attribute, bool(args[field]))
                if args.get("strike") is not None:
                    run.font.strike = bool(args["strike"])
                size = args.get("font_size") or args.get("size")
                if size:
                    run.font.size = Pt(float(size))
                colour = args.get("font_color") or args.get("color")
                if colour:
                    run.font.color.rgb = RGBColor.from_string(str(colour).lstrip("#"))
                hits += 1
        document.save(path)
        return result(f"formatted {hits} run(s) matching {find!r}",
                      warn=(f"0 runs matched {find!r}" if hits == 0 else None))

    if op == "format_paragraph":
        find = args.get("find")
        if not find:
            return "[error] format_paragraph needs 'find' (the anchor text)."
        hits = 0
        for paragraph in _paragraphs(document):
            if find not in paragraph.text:
                continue
            _spacing(paragraph, args)
            _align(paragraph, args.get("align"))
            _flow(paragraph, args)
            hits += 1
        document.save(path)
        return result(f"formatted {hits} paragraph(s) matching {find!r}",
                      warn=(f"0 matched {find!r}" if hits == 0 else None))

    if op == "add_hyperlink":
        find, url = args.get("find"), args.get("url")
        if not find or not url:
            return "[error] add_hyperlink needs 'find' and 'url'."
        for paragraph in _paragraphs(document):
            if find not in paragraph.text:
                continue
            from copy import deepcopy

            from docx.text.run import Run
            runs = [Run(node, paragraph) for node in paragraph._p.xpath("./w:r|./w:hyperlink/w:r")]
            whole = "".join(run.text for run in runs)
            at = whole.find(find)
            if at < 0:
                continue
            end, offset, matches = at + len(find), 0, []
            for run in runs:
                length = len(run.text)
                if offset < end and offset + length > at:
                    matches.append((run, offset))
                offset += length
            if any(run._r.getparent() is not paragraph._p for run, _ in matches):
                return "[error] add_hyperlink overlaps an existing hyperlink; existing content left unchanged"
            # Unlike the upstream whole-paragraph rewrite, leave every unrelated
            # run and hyperlink intact. Only the selected text becomes a new link.
            link = _hyperlink(paragraph, url, find)
            for index, (run, offset) in enumerate(matches):
                text = run.text
                tail = text[max(0, end - offset):]
                if index == 0:
                    run.text = text[:max(0, at - offset)]
                    run._r.addnext(link)
                    if tail:
                        node = deepcopy(run._r)
                        Run(node, paragraph).text = tail
                        link.addnext(node)
                else:
                    run.text = tail
            document.save(path)
            return f"added hyperlink on {find!r} → {url}"
        return result("add_hyperlink skipped", warn=f"text not found: {find!r}")

    if op == "set_page_margins":
        section = document.sections[int(args.get("section", 0))]
        for side in ("top", "bottom", "left", "right"):
            if args.get(side) is not None:
                setattr(section, f"{side}_margin", Inches(float(args[side])))
        document.save(path)
        return f"set page margins on section {args.get('section', 0)}"

    if op == "set_page_orientation":
        from docx.enum.section import WD_ORIENT
        section = document.sections[int(args.get("section", 0))]
        wanted = str(args.get("orientation", "portrait")).lower()
        landscape = wanted.startswith("land")
        section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        # Setting the orientation does not swap the page dimensions, and a landscape
        # section on a portrait-sized page is still a portrait page.
        width, height = section.page_width, section.page_height
        if (landscape and width < height) or (not landscape and width > height):
            section.page_width, section.page_height = height, width
        document.save(path)
        return f"set orientation={wanted} on section {args.get('section', 0)}"

    if op == "set_header_footer":
        section = document.sections[int(args.get("section", 0))]
        for name in ("header", "footer"):
            if args.get(name) is None:
                continue
            container = getattr(section, name)
            container.is_linked_to_previous = False
            paragraph = (container.paragraphs[0] if container.paragraphs
                         else container.add_paragraph())
            paragraph.text = args[name]
        document.save(path)
        return f"set header/footer on section {args.get('section', 0)}"

    if op == "set_page_number":
        summary = _page_number(document, args)
        document.save(path)
        return summary

    return f"[error] unknown docx op {op!r}: this writes {', '.join(OPS)}."
