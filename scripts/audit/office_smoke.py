"""Real Office acceptance: PYTHONPATH=. python scripts/audit/office_smoke.py OUTPUT [--render].

Requires LibreOffice and pdftotext; --render additionally needs pdftoppm. Uses only newly created
fixtures under OUTPUT. Unit-test mocks never count as a successful smoke run.
"""
from __future__ import annotations

import argparse
import asyncio
import hashlib
import json
import shutil
import subprocess
from pathlib import Path

import openpyxl
import pypdfium2
from pptx import Presentation
from pptx.dml.color import RGBColor

from misaka.core.documents.office import soffice
from misaka.core.tools.office import create_office_tool_definition
from misaka.core.tools.read import create_read_tool_definition


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


async def smoke(out, render):
    executable = soffice.binary()
    if executable is None:
        raise RuntimeError(soffice.INSTALL_HINT)
    version = (await asyncio.to_thread(subprocess.run, [executable, '--version'], check=True,
                                      capture_output=True, text=True, timeout=30)).stdout.strip()
    renderer = shutil.which('pdftoppm') if render else None
    text_reader = shutil.which('pdftotext')
    if text_reader is None:
        raise RuntimeError('CJK content acceptance requires pdftotext (Poppler)')
    if render and renderer is None:
        raise RuntimeError('--render requires pdftoppm')
    out.mkdir(parents=True, exist_ok=False)
    tool = create_office_tool_definition(str(out))
    read = create_read_tool_definition(str(out))
    records = {'libreoffice': version, 'formats': []}

    async def write(name, ops):
        result = await tool.execute('office-smoke', {'path': name, 'ops': ops})
        return result.content[0].text

    for suffix, args, expected in [
        ('docx', {'blocks': [{'type': 'heading', 'text': 'Office 中文验收'},
                            {'type': 'paragraph', 'text': 'Word content round trip.'},
                            {'type': 'table', 'rows': [['项目', '数值'], ['验证', 42]]}]}, 'Word content'),
        ('xlsx', {'sheets': [{'name': 'Values', 'headers': ['输入', '结果', '文本'],
                             'rows': [[2, '=SUM(A2,3)', 'placeholder']]}]}, 'Values'),
        ('pptx', {'slides': [{'layout': 'title_only', 'title': 'Office 中文验收',
                             'notes': 'Speaker notes survive duplication.'}]}, 'Office'),
    ]:
        name = 'sample.' + suffix
        path = out / name
        ops = [{'create': args}]
        if suffix == 'xlsx':
            ops += [{'set_cell': {'sheet': 'Values', 'cell': 'C2', 'value': '=literal', 'type': 'text'}},
                    {'set_cell': {'sheet': 'Values', 'cell': 'D2', 'value': '=IF(1=1,"",1)'}}]
        await write(name, ops)
        if suffix == 'pptx':
            await write(name, [{'add_chart': {'slide': 1, 'categories': ['A', 'B'],
                                             'series': {'Values': [1, 2]}, 'y': 1.5}}])
            deck = Presentation(path)
            deck.slides[0].background.fill.solid()
            deck.slides[0].background.fill.fore_color.rgb = RGBColor(240, 244, 248)
            deck.save(path)
            await write(name, [{'duplicate_slide': {'slide': 1}},
                               {'set_text': {'slide': 2, 'placeholder': 'title', 'text': 'Office 副本'}}])
            deck = Presentation(path)
            assert all(s.notes_slide.notes_text_frame.text == 'Speaker notes survive duplication.'
                       for s in deck.slides)
            charts = [next(shape.chart for shape in s.shapes if shape.has_chart) for s in deck.slides]
            assert charts[0].part is not charts[1].part
            assert charts[0].part.chart_workbook.xlsx_part is not charts[1].part.chart_workbook.xlsx_part
        if suffix == 'xlsx':
            book = openpyxl.load_workbook(path, data_only=True)
            assert book['Values']['B2'].value == 5
            assert book['Values']['C2'].value == '=literal'
            book.close()
            await write(name, [{'set_cell': {'sheet': 'Values', 'cell': 'A2', 'value': 7}}])
            book = openpyxl.load_workbook(path, data_only=True)
            assert book['Values']['B2'].value == 10
            book.close()
            receipt = await write(name, [{'set_cell': {'sheet': 'Values', 'cell': 'B3', 'value': '=1/0'}}])
            assert 'Values!B3 #DIV/0!' in receipt
            records['formula_results'] = {'initial': 5, 'after_input_edit': 10, 'error_receipt': receipt}
        before = digest(path)
        text = (await read.execute('read-smoke', {'path': name})).content[0].text
        assert expected in text, text
        assert digest(path) == before
        pdf = out / (suffix + '.pdf')
        receipt = await write(name, [{'export_pdf': {'out': str(pdf)}}])
        assert digest(path) == before
        document = pypdfium2.PdfDocument(str(pdf))
        texts = []
        for index in range(len(document)):
            page = document[index]
            textpage = page.get_textpage()
            texts.append(textpage.get_text_bounded())
            textpage.close(); page.close()
        document.close()
        assert all(t.strip() for t in texts)
        assert len(texts) == (2 if suffix == 'pptx' else 1)
        required = {'docx': ['中文验收', '项目', '数值', '验证'],
                    'xlsx': ['输入', '结果', '文本', '=literal'],
                    'pptx': ['中文验收', '副本']}[suffix]
        # PDFium can return CJK fallback-font runs in content-stream rather than visual
        # order. Keep that output as evidence; use Poppler's spatial order for phrases.
        layout = (await asyncio.to_thread(
            subprocess.run, [text_reader, '-layout', str(pdf), '-'],
            check=True, capture_output=True, text=True, timeout=60)).stdout
        assert all(word in layout for word in required), layout
        if renderer:
            await asyncio.to_thread(
                subprocess.run, [renderer, '-scale-to', '1400', '-png', str(pdf), str(out / suffix)],
                check=True, capture_output=True, timeout=60)
        legacy = {'docx': 'doc', 'xlsx': 'xls', 'pptx': 'ppt'}[suffix]
        meta = {}
        converted = soffice.convert(path, legacy, into=out / 'legacy', meta=meta)
        assert converted is not None, meta
        old = Path(converted); before = digest(old)
        recovered = (await read.execute('legacy-smoke', {'path': str(old)})).content[0].text
        assert ('输入' if suffix == 'xlsx' else expected) in recovered, recovered
        assert digest(old) == before
        records['formats'].append({'format': suffix, 'receipt': receipt, 'pdfium_text': texts,
                                   'pdf_layout_text': layout,
                                   'legacy': legacy, 'source_unchanged_on_read': True})
    (out / 'results.json').write_text(json.dumps(records, ensure_ascii=False, indent=2))
    print(json.dumps(records, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('output', type=Path)
    parser.add_argument('--render', action='store_true')
    arguments = parser.parse_args()
    asyncio.run(smoke(arguments.output.resolve(), arguments.render))
