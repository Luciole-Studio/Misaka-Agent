"""Differential Office smoke against a separately supplied, pinned upstream checkout.

Usage: PYTHONPATH=. .venv/bin/python scripts/audit/office_upstream.py UPSTREAM OUTDIR
No network, installs, real user documents, or upstream sandbox entry point is used.
"""
import ast
import hashlib
import json
import runpy
import subprocess
import sys
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from PIL import Image

from misaka.core.tools._office import docx, pptx, text, xlsx
from misaka.core.tools._office._receipt import normalise

PIN = "9e533db6f6c34d16037ee5ec964c479d0eb51cde"
ROOT = Path(__file__).resolve().parents[2]


def package(path):
    if path.suffix == ".txt":
        return {"text": path.read_text()}
    result = {}
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            data = archive.read(name)
            if name.endswith((".xml", ".rels")):
                tree = ET.fromstring(data)
                if name == "docProps/core.xml":
                    for child in list(tree):
                        if child.tag.endswith(("}created", "}modified")):
                            tree.remove(child)
                data = ET.canonicalize(ET.tostring(tree, encoding="unicode")).encode()
            result[name] = hashlib.sha256(data).hexdigest()
    return result


def main():
    upstream, out = map(Path, sys.argv[1:])
    actual = subprocess.check_output(["git", "-C", str(upstream), "rev-parse", "HEAD"], text=True).strip()
    if actual != PIN:
        raise ValueError(f"Expected upstream {PIN}, got {actual}")
    manifest = json.loads((ROOT / "misaka/core/tools/_office/PROVENANCE.json").read_text())
    for name, expected in manifest["upstream_files"].items():
        if hashlib.sha256((upstream / "plugins/tools" / name).read_bytes()).hexdigest() != expected:
            raise ValueError(f"Upstream source differs from pin: {name}")
    out.mkdir(parents=True, exist_ok=False)
    img = out / "image.png"
    Image.new("RGB", (2, 2)).save(img)
    programs = runpy.run_path(str(ROOT / "tests/office/test_office_operation_coverage.py"))["programs"](str(img.resolve()))
    report = {"upstream": PIN, "operations": [], "sources": {}}
    for source in sorted((upstream / "plugins/tools").glob("_*.py")):
        if source.stem.startswith(("_writer", "_reader", "_doc_reader")):
            report["sources"][source.name] = {
                "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                "functions": [node.name for node in ast.parse(source.read_text()).body
                              if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))]}
    for suffix, writer in [("docx", docx), ("xlsx", xlsx), ("pptx", pptx), ("txt", text)]:
        fmt = "text" if suffix == "txt" else suffix
        context = {"__name__": "office_audit_upstream"}
        # The upstream build concatenates the same core + format modules. Do not run
        # main(), which invokes sandbox-specific paths or recalculation macros.
        for name in ("_writer_core", f"_writer_{fmt}"):
            source = upstream / "plugins/tools" / f"{name}.py"
            exec(compile(source.read_text(), str(source), "exec"), context)  # noqa: S102 - explicit pinned audit fixture
        context["_ensure"] = lambda module, _package: __import__(module)
        original = context[f"_{fmt}_write"]
        ours, theirs = out / f"misaka.{suffix}", out / f"frontier.{suffix}"
        for op, args in programs[suffix]:
            our_result = normalise(op, writer.write(str(ours), op, args))
            their_result = context["_norm_result"](op, original(str(theirs), op, args))
            if not our_result["ok"] or not their_result["ok"]:
                raise AssertionError((suffix, op, our_result, their_result))
            a, b = package(ours), package(theirs)
            changed = sorted(key for key in a.keys() | b.keys() if a.get(key) != b.get(key))
            report["operations"].append({"format": suffix, "op": op, "both_ok": True,
                                         "different_package_parts": changed})
    # Compare extraction primitives on identical documents, separately from the
    # deliberate quote-safe presentation of rendered text.
    from openpyxl import load_workbook
    from pptx import Presentation

    from misaka.core.documents.office import pptx as pr
    from misaka.core.documents.office import xlsx as xr
    report["reader_checks"] = []
    def compare(fmt, upstream_name, local_name, *args):
        ours = getattr(xr if fmt == "xlsx" else pr, local_name)(*args)
        theirs = readers[fmt][upstream_name](*args)
        check = {"format": fmt, "function": upstream_name, "equal": ours == theirs}
        if ours != theirs and upstream_name == "_x_pivot_regions":
            check["difference"] = "MISAKA appends the explicit sheet/range to the re-read hint"
            check["misaka"] = repr(ours[0])
            check["upstream"] = repr(theirs[0])
        report["reader_checks"].append(check)
    readers = {}
    for fmt in ("xlsx", "pptx"):
        context = {"__name__": "office_reader_audit"}
        for name in ("_reader_core", f"_reader_{fmt}"):
            source = upstream / "plugins/tools" / f"{name}.py"
            exec(compile(source.read_text(), str(source), "exec"), context)  # noqa: S102 - pinned fixture
        readers[fmt] = context
    book = load_workbook(out / "misaka.xlsx")
    sheet = book.active
    sheet["H1"] = "=SUM(B2:C2)"
    sheet["H2"] = "=SUM(B3:C3)"
    sheet.merge_cells("K1:L1")
    from openpyxl.chart import BarChart, Reference
    from openpyxl.pivot.cache import CacheDefinition, CacheSource, WorksheetSource
    from openpyxl.pivot.table import Location, TableDefinition
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, max_col=3, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
    sheet.add_chart(chart, "J5")
    pivot = TableDefinition(name="P", cacheId=1, dataCaption="Data", location=Location(
        ref="M1:N3", firstHeaderRow=1, firstDataRow=1, firstDataCol=1))
    pivot.cache = CacheDefinition(cacheSource=CacheSource(type="worksheet", worksheetSource=WorksheetSource(ref="A1:C3", sheet="S")))
    sheet.add_pivot(pivot)
    book.save(out / "reader-fixture.xlsx")
    coords = {(cell.row, cell.column) for row in sheet for cell in row if cell.value is not None}
    for upstream_name, local_name, args in [
        ("_x_col", "_col", (28,)), ("_x_ref", "_ref", (1, 1, 3, 4)),
        ("_x_islands", "_islands", (coords,)), ("_x_compress", "_compress", (coords,)),
        ("_x_r1c1", "_r1c1", ("=SUM(A1,$B$2,C$3)", 3, 4)),
        ("_x_formula_lines", "_formula_lines", (sheet,)),
        ("_x_numfmt_lines", "_numfmt_lines", (sheet, coords)),
        ("_x_merged_lines", "_merged_lines", (sheet,)),
        ("_x_style_lines", "_style_lines", (sheet, coords)),
        ("_x_cond_lines", "_cond_lines", (sheet,)),
        ("_x_extra_lines", "_extra_lines", (sheet,)),
        ("_x_pivot_regions", "_pivot_regions", (sheet,)),
        ("_x_chart_lines", "_chart_lines", (str(out / "reader-fixture.xlsx"),)),
    ]:
        compare("xlsx", upstream_name, local_name, *args)
    for value in (0, 12.5, -2, "a", "", "1,234.5", "20%"):
        compare("xlsx", "_x_to_num", "_to_num", value)
    book.close()
    deck = Presentation(out / "misaka.pptx")
    for shape in deck.slides[0].shapes:
        for upstream_name, local_name in [("_ph_type", "_placeholder"), ("_shape_fill", "_fill"),
                                          ("_shape_bits", "_bits"), ("_bbox", "_bbox"),
                                          ("_is_line", "_is_line"), ("_is_box", "_is_box")]:
            compare("pptx", upstream_name, local_name, shape)
    (out / "comparison.json").write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps({"operations": len(report["operations"]), "package_equal": sum(
        not item["different_package_parts"] for item in report["operations"]),
        "reader_checks": len(report["reader_checks"]),
        "reader_equal": sum(item["equal"] for item in report["reader_checks"]),
        "report": str(out / "comparison.json")}))


if __name__ == "__main__":
    main()
