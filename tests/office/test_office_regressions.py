"""Integration regressions found against FrontierAgent 9e533db6 (2026-09-15)."""
import json
from pathlib import Path

import pytest

from misaka.core.documents.office import soffice
from misaka.core.tools import _office
from misaka.core.tools._office import _intent
from misaka.core.tools.office import create_office_tool_definition


@pytest.fixture
def pdf_converter(monkeypatch):
    monkeypatch.setattr(soffice, "binary", lambda: "soffice")
    def export(source, out, **kwargs):
        Path(out).write_bytes(b"%PDF-1.7\n" + Path(source).read_bytes())
        return True
    monkeypatch.setattr(soffice, "export_pdf", export)


async def test_failed_batch_is_a_tool_error(tmp_path):
    tool = create_office_tool_definition(str(tmp_path))
    with pytest.raises(RuntimeError, match="STOPPED"):
        await tool.execute("call", {"path": "x.docx", "ops": [{"teleport": {}}]})
    assert not (tmp_path / "x.docx").exists()


def test_export_default_name_is_target_name(tmp_path, pdf_converter):
    path = tmp_path / "report.docx"
    receipt = _office.run_ops(path, [{"create": {"blocks": [{"type": "paragraph", "text": "report"}]}}, {"export_pdf": {}}])
    assert receipt.startswith("✓"), receipt
    assert (tmp_path / "report.pdf").read_bytes().startswith(b"%PDF")
    assert not list(tmp_path.glob(".office-*"))


@pytest.mark.parametrize("existing", [False, True])
def test_export_rolls_back_with_failed_batch(tmp_path, pdf_converter, existing):
    path, pdf = tmp_path / "report.docx", tmp_path / "report.pdf"
    if existing:
        pdf.write_bytes(b"original pdf")
    receipt = _office.run_ops(path, [{"create": {"blocks": [{"type": "paragraph", "text": "report"}]}},
        {"export_pdf": {"out": str(pdf)}}, {"teleport": {}}])
    assert "STOPPED" in receipt
    assert not path.exists()
    assert pdf.read_bytes() == b"original pdf" if existing else not pdf.exists()
    assert "file unchanged" in receipt


def test_publish_failure_is_not_masked_by_index_error(tmp_path, monkeypatch):
    def fail(*args):
        raise PermissionError("publish denied")
    monkeypatch.setattr(_office.os, "replace", fail)
    receipt = _office.run_ops(tmp_path / "x.txt", [{"create": {"content": "x"}}])
    assert "publish denied" in receipt
    assert "IndexError" not in receipt
    assert not (tmp_path / "x.txt").exists()


def test_intent_sequence_survives_four_digits(tmp_path):
    target = tmp_path / "x.txt"
    bucket = Path(_intent._bucket(target, workspace=tmp_path))
    bucket.mkdir(parents=True)
    (bucket / "999_create.json").write_text("{}")
    (bucket / "1000_append.json").write_text("{}")
    saved = _intent.archive(target, [{"append": {"content": "new"}}], workspace=tmp_path)
    assert json.loads(Path(saved).read_text())["seq"] == 1001
    assert (bucket / "1000_append.json").read_text() == "{}"


def test_recalc_zero_exit_without_new_output_is_not_success(tmp_path, monkeypatch):
    import subprocess
    source = tmp_path / "source.xlsx"
    source.write_bytes(b"not recalculated")
    monkeypatch.setattr(soffice, "binary", lambda: "soffice")
    monkeypatch.setattr(soffice.subprocess, "run", lambda *a, **k: subprocess.CompletedProcess(a, 0, "", ""))
    (tmp_path / "out").mkdir()
    meta = {}
    assert soffice.recalc(source, into=tmp_path / "out", meta=meta) is None
    assert "no xlsx produced" in meta["soffice_error"]


async def test_relative_image_paths_use_tool_workspace(tmp_path):
    from PIL import Image
    Image.new("RGB", (2, 2)).save(tmp_path / "image.png")
    tool = create_office_tool_definition(str(tmp_path))
    receipt = await tool.execute("call", {"path": "image.docx", "ops": [
        {"create": {"blocks": [{"type": "image", "path": "image.png"}]}}]})
    assert receipt.content[0].text.startswith("✓")
    from docx import Document
    assert len(Document(tmp_path / "image.docx").inline_shapes) == 1


def test_secondary_export_obeys_subagent_workspace_guard(tmp_path):
    from misaka.core.subagent.policy import _permission_restriction
    args = {"path": "report.docx", "ops": [{"export_pdf": {"out": str(tmp_path.parent / "other.pdf")}}]}
    assert "outside" in _permission_restriction([], "office", args, str(tmp_path))
    args["ops"][0]["export_pdf"]["out"] = str(tmp_path / "report.pdf")
    assert _permission_restriction([], "office", args, str(tmp_path)) is None
    (tmp_path / "batch.json").write_text(json.dumps([{"export_pdf": {"out": str(tmp_path.parent / "other.pdf")}}]))
    args["ops"] = "@batch.json"
    assert "outside" in _permission_restriction([], "office", args, str(tmp_path))


def test_json_content_path_fields_remain_literal(tmp_path):
    from misaka.core.tools._office.paths import resolve_ops
    ops = [{"create": {"data": {"path": "relative", "out": "literal"}}}]
    assert resolve_ops(ops, str(tmp_path)) == ops


@pytest.mark.parametrize("suffix", ["docm", "xlsm", "pptm"])
def test_macro_writes_are_not_silently_mislabelled(tmp_path, suffix):
    receipt = _office.run_ops(tmp_path / f"x.{suffix}", [{"create": {}}])
    assert receipt.startswith("[error]")
    assert not (tmp_path / f"x.{suffix}").exists()


def test_legacy_raw_rows_are_readable():
    from misaka.core.documents.office.xlsx import render_rows
    rendered = render_rows("S", [["year", "amount"], [2026, 12.5]])
    assert "2026\t12.5" in rendered


async def test_read_legacy_file_uses_conversion(tmp_path, monkeypatch):
    from docx import Document

    from misaka.core.tools.read import create_read_tool_definition
    path = tmp_path / "old.doc"
    path.write_bytes(b"legacy")
    before = path.read_bytes()
    def convert(source, target_suffix, *, into, **kwargs):
        target = Path(into, "old.docx")
        doc = Document()
        doc.add_paragraph("converted legacy content")
        doc.save(target)
        return str(target)
    monkeypatch.setattr(soffice, "binary", lambda: "soffice")
    monkeypatch.setattr(soffice, "convert", convert)
    result = await create_read_tool_definition(str(tmp_path)).execute("read", {"path": "old.doc"})
    assert "converted legacy content" in result.content[0].text
    assert path.read_bytes() == before


def test_chart_categories_in_middle_do_not_drop_data(tmp_path):
    from openpyxl import load_workbook
    path = tmp_path / "x.xlsx"
    receipt = _office.run_ops(path, [{"create": {"sheets": [{"name": "S", "rows": [
        ["first", "category", "last"], [1, "A", 3], [2, "B", 4]]}]}},
        {"add_chart": {"sheet": "S", "data_range": "A1:C3", "categories_col": "B"}}])
    assert receipt.startswith("✓"), receipt
    chart = load_workbook(path).active._charts[0]
    assert [series.val.numRef.f for series in chart.series] == ["'S'!$A$2:$A$3", "'S'!$C$2:$C$3"]


def test_recalculation_reports_errors_and_runs_after_input_edit(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    path = tmp_path / "x.xlsx"
    _office.run_ops(path, [{"create": {"sheets": [{"name": "S", "rows": [[1, "=1/A1"]]}]}}])
    def recalc(source, meta):
        book = load_workbook(source)
        book.active["B1"] = "#DIV/0!"
        book.save(source)
        return True
    monkeypatch.setattr(_office, "_recalculate", recalc)
    receipt = _office.run_ops(path, [{"set_cell": {"cell": "A1", "value": 0}}])
    assert "1 error(s): S!B1 #DIV/0!" in receipt
    assert "fidelity best-effort" in receipt


def test_delete_then_add_slide_has_unique_zip_members(tmp_path):
    import zipfile

    from pptx import Presentation
    path = tmp_path / "deck.pptx"
    receipt = _office.run_ops(path, [{"create": {"slides": [{"title": "one"}, {"title": "two"}]}},
        {"delete_slide": {"slide": 1}}, {"add_slide": {"title": "three"}}])
    assert receipt.startswith("✓"), receipt
    assert [slide.shapes.title.text for slide in Presentation(path).slides] == ["two", "three"]
    with zipfile.ZipFile(path) as archive:
        assert len(archive.namelist()) == len(set(archive.namelist()))


async def test_cancelled_save_retains_mutation_lock_until_thread_finishes(tmp_path, monkeypatch):
    import asyncio
    import threading

    from misaka.core.tools.file_mutation_queue import with_file_mutation_queue
    entered, release = threading.Event(), threading.Event()
    real_run = _office.run_ops
    def slow_run(*args, **kwargs):
        entered.set()
        assert release.wait(5)
        return real_run(*args, **kwargs)
    monkeypatch.setattr(_office, "run_ops", slow_run)
    path = tmp_path / "x.txt"
    call = asyncio.create_task(create_office_tool_definition(str(tmp_path)).execute(
        "call", {"path": str(path), "content": "saved"}))
    second_entered = asyncio.Event()
    async def second():
        second_entered.set()
    other = None
    try:
        assert await asyncio.to_thread(entered.wait, 5)
        call.cancel()
        other = asyncio.create_task(with_file_mutation_queue(str(path), second))
        await asyncio.sleep(0.02)
        assert not second_entered.is_set()
        release.set()
        with pytest.raises(asyncio.CancelledError):
            await call
        await other
        assert path.read_text() == "saved"
    finally:
        release.set()
        await asyncio.gather(call, *([other] if other else []), return_exceptions=True)


def test_publish_second_failure_restores_all_destinations(tmp_path, pdf_converter, monkeypatch):
    path, pdf = tmp_path / "x.docx", tmp_path / "x.pdf"
    _office.run_ops(path, [{"create": {"blocks": [{"text": "old"}]}}])
    before = path.read_bytes()
    pdf.write_bytes(b"old pdf")
    replace = _office.os.replace
    def fail_second(source, target):
        if Path(target) == pdf and str(source).endswith(".pdf"):
            raise PermissionError("pdf publish denied")
        return replace(source, target)
    monkeypatch.setattr(_office.os, "replace", fail_second)
    receipt = _office.run_ops(path, [{"insert_paragraph": {"text": "new"}}, {"export_pdf": {}}])
    assert "pdf publish denied" in receipt
    assert path.read_bytes() == before
    assert pdf.read_bytes() == b"old pdf"
    assert not list(tmp_path.glob(".office-*"))


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16"])
def test_entities_after_long_xml_prolog_are_rejected(tmp_path, encoding):
    import zipfile

    from misaka.core.documents.office import precheck
    path = tmp_path / "entities.docx"
    payload = ('<?xml version="1.0" encoding="'+encoding+'"?>\n' + " " * 70000
               + '<!DOCTYPE x [<!ENTITY injected "value">]><x>&injected;</x>')
    with zipfile.ZipFile(path, "w") as package:
        package.writestr("[Content_Types].xml", "<Types/>")
        package.writestr("word/document.xml", payload.encode(encoding))
    with pytest.raises(ValueError, match="entities"):
        precheck(path)


def test_workbook_reader_does_not_modify_global_pivot_parser(tmp_path):
    from openpyxl import Workbook
    from openpyxl.reader.workbook import WorkbookParser

    from misaka.core.documents.office import xlsx as reader
    before = WorkbookParser.pivot_caches
    book = Workbook()
    book.active["A1"] = 1
    path = tmp_path / "x.xlsx"
    book.save(path)
    reader.render(str(path))
    assert WorkbookParser.pivot_caches is before


def test_docm_and_content_controls_are_read_without_dropping_text(tmp_path):
    import zipfile

    from docx import Document
    from docx.oxml import OxmlElement

    from misaka.core.documents import office
    path, macro = tmp_path / "x.docx", tmp_path / "x.docm"
    doc = Document()
    para = doc.add_paragraph("inside content control")
    sdt, content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
    para._p.addprevious(sdt)
    sdt.append(content)
    content.append(para._p)
    doc.save(path)
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(macro, "w") as output:
        for part in source.infolist():
            data = source.read(part)
            if part.filename == "[Content_Types].xml":
                data = data.replace(b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
                                    b"application/vnd.ms-word.document.macroEnabled.main+xml")
            output.writestr(part, data)
    before = macro.read_bytes()
    assert "inside content control" in office.render(str(macro))
    assert macro.read_bytes() == before


def test_pptx_richtext_strike_is_not_silently_ignored(tmp_path):
    from pptx import Presentation
    path = tmp_path / "x.pptx"
    _office.run_ops(path, [{"create": {"slides": [{"title": [{"text": "old", "strike": True}]}]}}])
    run = Presentation(path).slides[0].shapes.title.text_frame.paragraphs[0].runs[0]
    assert run._r.rPr.get("strike") == "sngStrike"


def test_word_replace_in_link_preserves_other_runs(tmp_path):
    from docx import Document
    path = tmp_path / "x.docx"
    _office.run_ops(path, [{"create": {"blocks": [{"type": "paragraph", "text": [
        {"text": "Lead", "bold": True}, {"text": " linked", "link": "https://example.invalid/"}, " tail"]}]}},
        {"replace_text": {"find": "linked", "replace": "updated"}}])
    paragraph = Document(path).paragraphs[0]
    assert paragraph.text == "Lead updated tail"
    assert paragraph.runs[0].bold is True
    assert paragraph.hyperlinks[0].text == " updated"


def test_write_after_reader_retains_real_pivot_cache(tmp_path):
    from openpyxl import Workbook, load_workbook
    from openpyxl.pivot.cache import CacheDefinition, CacheSource, WorksheetSource
    from openpyxl.pivot.table import Location, TableDefinition

    from misaka.core.documents.office import xlsx as reader
    path = tmp_path / "pivot.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "S"
    sheet.append(["label", "value"])
    sheet.append(["one", 1])
    pivot = TableDefinition(name="P", cacheId=1, dataCaption="Data", location=Location(
        ref="D1:E3", firstHeaderRow=1, firstDataRow=1, firstDataCol=1))
    pivot.cache = CacheDefinition(cacheSource=CacheSource(type="worksheet", worksheetSource=WorksheetSource(ref="A1:B2", sheet="S")))
    sheet.add_pivot(pivot)
    book.save(path)
    rendered = reader.render(str(path))
    assert "source=S!A1:B2" in rendered
    receipt = _office.run_ops(path, [{"set_cell": {"cell": "B2", "value": 2}}])
    assert receipt.startswith("✓"), receipt
    loaded = load_workbook(path)
    assert loaded.active._pivots[0].cache.cacheSource.worksheetSource.ref == "A1:B2"


def test_adding_hyperlink_preserves_surrounding_format_and_other_links(tmp_path):
    from docx import Document
    path = tmp_path / "x.docx"
    _office.run_ops(path, [{"create": {"blocks": [{"text": [
        {"text": "Bold lead ", "bold": True}, "target then ",
        {"text": "other link", "link": "https://example.invalid/old"},
        {"text": " tail", "italic": True}]}]}},
        {"add_hyperlink": {"find": "target", "url": "https://example.invalid/new"}}])
    paragraph = Document(path).paragraphs[0]
    assert paragraph.text == "Bold lead target then other link tail"
    assert paragraph.runs[0].bold is True
    assert paragraph.runs[-1].italic is True
    assert [link.text for link in paragraph.hyperlinks] == ["target", "other link"]
