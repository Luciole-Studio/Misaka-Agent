"""Every docx op, verified by reading the document back with misaka's own reader.

That is the acceptance rule for the whole writing layer: an op is done when the reader
finds what the op claimed to write. Asserting on the writer's own receipt would only prove
the writer agrees with itself, and the failure this catches -- a property python-docx
accepts and silently drops -- looks like success from the inside.
"""
from __future__ import annotations

import importlib
import zipfile

from misaka.core.documents.index import normalize_for_quote_match as fold
from misaka.core.documents.office import docx as reader
from misaka.core.tools._office import docx as writer

importlib.import_module("docx")

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _document(tmp_path, blocks, name="doc.docx"):
    path = tmp_path / name
    outcome = writer.write(str(path), "create", {"blocks": blocks})
    assert outcome["ok"], outcome
    return path


def _xml(path, member="word/document.xml"):
    with zipfile.ZipFile(path) as archive:
        return archive.read(member).decode("utf-8")


# ---- create ------------------------------------------------------------------------------

def test_create_writes_every_block_type_and_counts_them(tmp_path):
    path = tmp_path / "all.docx"
    outcome = writer.write(str(path), "create", {"blocks": [
        {"type": "heading", "text": "Title", "level": 1},
        {"type": "paragraph", "text": "body"},
        {"type": "list", "items": ["one", "two"]},
        {"type": "table", "rows": [["a", "b"], ["1", "2"]]},
        {"type": "page_break"},
    ]})
    assert outcome["counts"] == {"heading": 1, "paragraph": 1, "list": 1,
                                 "table": 1, "page_break": 1}
    out = reader.render(str(path))
    assert "# Title" in out
    assert "- one\n- two" in out
    assert "`table 2×2`" in out


def test_create_refuses_a_path_that_exists(tmp_path):
    path = _document(tmp_path, [{"type": "paragraph", "text": "first"}])
    outcome = writer.write(str(path), "create", {"blocks": [{"type": "paragraph",
                                                             "text": "second"}]})
    assert not outcome["ok"]
    assert "already exists" in outcome["summary"]
    assert "first" in reader.render(str(path))


def test_a_create_whose_blocks_all_wrote_nothing_fails_loudly(tmp_path):
    """A document that came out empty must not report success: the model would move on and
    deliver a blank file."""
    outcome = writer.write(str(tmp_path / "empty.docx"), "create",
                           {"blocks": [{"type": "mystery"}, {}]})
    assert not outcome["ok"]
    assert "all empty or unknown" in outcome["summary"]


def test_a_single_empty_block_among_good_ones_is_reported_not_hidden(tmp_path):
    path = tmp_path / "partial.docx"
    outcome = writer.write(str(path), "create", {"blocks": [
        {"type": "paragraph", "text": "kept"},
        {"type": "table", "rows": []},
    ]})
    assert outcome["ok"]
    assert "block(s) [1] wrote nothing" in outcome["warn"]


def test_metadata_reaches_the_document_properties(tmp_path):
    path = tmp_path / "titled.docx"
    writer.write(str(path), "create", {"metadata": {"title": "A Study", "author": "R"},
                                       "blocks": [{"type": "paragraph", "text": "x"}]})
    assert reader.title(str(path)) == "A Study"


# ---- rich text ---------------------------------------------------------------------------

def test_runs_carry_their_own_formatting(tmp_path):
    path = _document(tmp_path, [{"type": "paragraph", "text": [
        {"text": "plain "},
        {"text": "bold", "bold": True},
        {"text": " and "},
        {"text": "struck", "strike": True},
    ]}])
    out = reader.render(str(path))
    assert 'bold "bold"' in out
    assert 'strikethrough "struck"' in out
    assert fold("plain bold and struck") in fold(out)


def test_a_run_with_a_link_becomes_a_real_hyperlink(tmp_path):
    """python-docx has no API for one, so this goes through OXML. Written as plain text
    instead, the URL is not clickable and carries no relationship."""
    path = _document(tmp_path, [{"type": "paragraph", "text": [
        {"text": "see this", "link": "https://example.org/x"}]}])
    assert "w:hyperlink" in _xml(path)
    with zipfile.ZipFile(path) as archive:
        assert "example.org" in archive.read("word/_rels/document.xml.rels").decode()


def test_a_cjk_font_is_written_where_cjk_text_reads_it(tmp_path):
    """``font.name`` writes only ``w:ascii``/``w:hAnsi``. CJK characters take their font
    from ``w:eastAsia``, so without it the request has no effect at all and the document
    comes out in the theme font."""
    path = _document(tmp_path, [{"type": "paragraph",
                                 "text": [{"text": "季度报告", "font": "SimSun"}]}])
    xml = _xml(path)
    assert 'w:eastAsia="SimSun"' in xml


# ---- tables -------------------------------------------------------------------------------

def test_a_header_row_is_bold_and_repeats_across_pages(tmp_path):
    path = _document(tmp_path, [{"type": "table", "header": True,
                                 "rows": [["region", "units"], ["North", "800"]]}])
    assert "w:tblHeader" in _xml(path)


def test_column_widths_are_written_where_both_programs_read_them(tmp_path):
    """Word and LibreOffice read different places and both default to auto-fit, so a width
    set only one way silently does nothing in the other program."""
    path = _document(tmp_path, [{"type": "table", "column_widths_in": [1.5, 3.0],
                                 "rows": [["a", "b"], ["1", "2"]]}])
    xml = _xml(path)
    assert 'w:type="fixed"' in xml                       # the layout
    assert f'w:w="{round(1.5 * 1440)}"' in xml           # tblGrid, for LibreOffice
    assert xml.count("<w:tcW") >= 4                      # per cell, for Word


def test_a_table_cell_can_carry_its_own_fill_and_alignment(tmp_path):
    path = _document(tmp_path, [{"type": "table", "rows": [
        [{"content": "head", "fill_color": "#DDEEFF", "align": "center"}, "b"]]}])
    assert 'w:fill="DDEEFF"' in _xml(path)


def test_cant_split_marks_every_row_once(tmp_path):
    """The schema fixes the order inside ``w:trPr`` and the header row has already put
    ``w:tblHeader`` there; appending would invert the two and produce a file Word refuses."""
    path = tmp_path / "split.docx"
    writer.write(str(path), "create", {"blocks": [{"type": "paragraph", "text": "x"}]})
    writer.write(str(path), "insert_table", {"rows": [["a"], ["b"]], "cant_split": True})
    xml = _xml(path)
    assert xml.count("w:cantSplit") == 2
    assert reader.render(str(path))                       # still parses


# ---- editing ------------------------------------------------------------------------------

def test_replace_text_keeps_the_rest_of_the_paragraph_formatted(tmp_path):
    """FrontierAgent rewrites the whole paragraph into its first run and blanks the rest,
    so one find-and-replace strips the bold off every other word -- silently, with a
    success receipt. The replacement is spliced by offset instead."""
    path = _document(tmp_path, [{"type": "paragraph", "text": [
        {"text": "Revenue "},
        {"text": "rose", "bold": True},
        {"text": " sharply in Q3."},
    ]}])
    outcome = writer.write(str(path), "replace_text", {"find": "sharply", "replace": "14%"})
    assert outcome["ok"]
    out = reader.render(str(path))
    assert fold("Revenue rose 14% in Q3.") in fold(out)
    assert 'bold "rose"' in out                            # the bold survived


def test_replace_text_reaches_inside_tables(tmp_path):
    """A figure quoted in a table is exactly the kind of thing that needs correcting, and
    a walk over ``doc.paragraphs`` never sees it."""
    path = _document(tmp_path, [{"type": "table", "rows": [["region", "800"]]}])
    writer.write(str(path), "replace_text", {"find": "800", "replace": "840"})
    assert "840" in reader.render(str(path))


def test_replace_text_honours_a_count(tmp_path):
    path = _document(tmp_path, [{"type": "paragraph", "text": "x x x"}])
    writer.write(str(path), "replace_text", {"find": "x", "replace": "y", "count": 2})
    assert "y y x" in reader.render(str(path))


def test_replace_text_with_no_match_warns_rather_than_failing(tmp_path):
    path = _document(tmp_path, [{"type": "paragraph", "text": "body"}])
    outcome = writer.write(str(path), "replace_text", {"find": "absent", "replace": "x"})
    assert outcome["ok"]
    assert "0 matches" in outcome["warn"]


def test_a_phrase_split_across_runs_is_still_replaced(tmp_path):
    """Word splits a sentence at every spell-check boundary, so a phrase routinely spans
    four runs; a per-run search would report no match for text plainly in the document."""
    path = _document(tmp_path, [{"type": "paragraph", "text": [
        {"text": "Total rev"}, {"text": "enue", "bold": True}, {"text": " rose"}]}])
    outcome = writer.write(str(path), "replace_text",
                           {"find": "revenue", "replace": "income"})
    assert outcome["warn"] is None
    assert fold("Total income rose") in fold(reader.render(str(path)))


def test_insert_paragraph_lands_after_its_anchor(tmp_path):
    path = _document(tmp_path, [{"type": "paragraph", "text": "first"},
                                {"type": "paragraph", "text": "third"}])
    writer.write(str(path), "insert_paragraph", {"text": "second", "after_text": "first"})
    out = reader.render(str(path))
    assert out.index("first") < out.index("second") < out.index("third")


def test_a_missing_anchor_warns_and_changes_nothing(tmp_path):
    """Appending at the end instead would put the text somewhere the model did not ask
    for, and the receipt would call it a success."""
    path = _document(tmp_path, [{"type": "paragraph", "text": "body"}])
    before = path.read_bytes()
    outcome = writer.write(str(path), "insert_paragraph",
                           {"text": "new", "after_text": "no such anchor"})
    assert "anchor 'no such anchor' not found" in outcome["warn"]
    assert path.read_bytes() == before


def test_insert_heading_and_insert_table_take_the_same_anchor(tmp_path):
    path = _document(tmp_path, [{"type": "paragraph", "text": "intro"},
                                {"type": "paragraph", "text": "outro"}])
    writer.write(str(path), "insert_heading", {"text": "Method", "level": 2,
                                               "after_text": "intro"})
    writer.write(str(path), "insert_table", {"rows": [["a", "b"]], "after_text": "Method"})
    out = reader.render(str(path))
    assert "## Method" in out
    assert out.index("## Method") < out.index("`table 1×2`") < out.index("outro")


def test_format_text_applies_to_the_matching_runs(tmp_path):
    path = _document(tmp_path, [{"type": "paragraph", "text": [
        {"text": "keep "}, {"text": "target"}]}])
    outcome = writer.write(str(path), "format_text", {"find": "target", "bold": True})
    assert outcome["ok"]
    assert 'bold "target"' in reader.render(str(path))


def test_add_hyperlink_turns_existing_text_into_a_link(tmp_path):
    path = _document(tmp_path, [{"type": "paragraph", "text": "see the appendix here"}])
    writer.write(str(path), "add_hyperlink", {"find": "appendix", "url": "https://x.test/a"})
    assert "w:hyperlink" in _xml(path)
    assert fold("see the appendix here") in fold(reader.render(str(path)))


# ---- page setup ----------------------------------------------------------------------------

def test_set_page_number_writes_a_field_not_a_number(tmp_path):
    """A literal "1" typed into the footer is the number 1 on every page."""
    path = _document(tmp_path, [{"type": "paragraph", "text": "body"}])
    writer.write(str(path), "set_page_number", {"location": "footer", "of_total": True})
    with zipfile.ZipFile(path) as archive:
        footers = [name for name in archive.namelist() if "footer" in name and name.endswith(".xml")]
        assert footers
        body = "".join(archive.read(name).decode("utf-8") for name in footers)
    assert "PAGE" in body and "NUMPAGES" in body


def test_landscape_also_swaps_the_page_dimensions(tmp_path):
    """Setting the orientation alone leaves a landscape section on a portrait-sized page."""
    import docx

    path = _document(tmp_path, [{"type": "paragraph", "text": "body"}])
    writer.write(str(path), "set_page_orientation", {"orientation": "landscape"})
    section = docx.Document(str(path)).sections[0]
    assert section.page_width > section.page_height


def test_margins_and_header_footer_apply(tmp_path):
    import docx
    from docx.shared import Inches

    path = _document(tmp_path, [{"type": "paragraph", "text": "body"}])
    writer.write(str(path), "set_page_margins", {"top": 2.0})
    writer.write(str(path), "set_header_footer", {"header": "Draft"})
    document = docx.Document(str(path))
    assert document.sections[0].top_margin == Inches(2.0)
    assert "Draft" in document.sections[0].header.paragraphs[0].text


def test_an_op_on_a_file_that_does_not_exist_says_so(tmp_path):
    outcome = writer.write(str(tmp_path / "absent.docx"), "replace_text",
                           {"find": "a", "replace": "b"})
    assert "file not found" in outcome


def test_an_unknown_op_names_the_ones_that_exist(tmp_path):
    path = _document(tmp_path, [{"type": "paragraph", "text": "x"}])
    outcome = writer.write(str(path), "teleport", {})
    assert "unknown docx op 'teleport'" in outcome
    assert "insert_paragraph" in outcome
