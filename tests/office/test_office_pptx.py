"""A deck reads as a diagram, not as a bag of strings, and stays quotable.

FrontierAgent's ``_reader_pptx.py`` is the only reader either project has that recovers
what a slide *means*: a chart's real numbers (they live in the embedded chart part, where
no text extraction reaches), the colour block and the caption drawn on top of it merged
into one line, connectors resolved into ``A -> B -> C``, and a row of identically styled
boxes hoisted into one group so the five labels that differ are not buried under five
identical style tags. All of that is ported here.

Three things are deliberately not ported, each with its own test below: the slide marker is
``## Slide N`` because ``office.paging`` cuts the deck there and the page number has to name
something; tables and chart rows are tab-separated; and inline formatting moves off the
line. The last two are one rule -- see ``test_office_docx``'s docstring for the folding
argument, which is the same here.
"""
from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt

from misaka.core.documents.index import normalize_for_quote_match as fold
from misaka.core.documents.office import pptx as reader


def _save(presentation, tmp_path, name="deck.pptx"):
    path = tmp_path / name
    presentation.save(path)
    return str(path)


def _deck(tmp_path):
    """Bullets, a table and a chart -- one slide each."""
    presentation = Presentation()

    first = presentation.slides.add_slide(presentation.slide_layouts[1])
    first.shapes.title.text = "Findings"
    frame = first.placeholders[1].text_frame
    frame.text = "First point"
    nested = frame.add_paragraph()
    nested.text = "Sub point"
    nested.level = 1
    first.notes_slide.notes_text_frame.text = "mention the caveat"

    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second.shapes.title.text = "Regions"
    table = second.shapes.add_table(
        3, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
    for row, (name, units) in enumerate(
            (("region", "units"), ("North", "800"), ("South", "440"))):
        table.cell(row, 0).text = name
        table.cell(row, 1).text = units

    third = presentation.slides.add_slide(presentation.slide_layouts[5])
    third.shapes.title.text = "Growth"
    data = CategoryChartData()
    data.categories = ["Q1", "Q2"]
    data.add_series("revenue", (1200.0, 1500.0))
    third.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                           Inches(1), Inches(2), Inches(6), Inches(4), data)
    return _save(presentation, tmp_path)


def _flow(tmp_path):
    """Four identically styled boxes joined by three connectors."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index, label in enumerate(("Collect", "Clean", "Analyse", "Report")):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.5 + index * 2.2), Inches(2), Inches(1.8), Inches(1))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0x22, 0x88, 0x99)
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = label
        run.font.size = Pt(18)
    for index in range(3):
        slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(2.3 + index * 2.2), Inches(2.5),
            Inches(2.7 + index * 2.2), Inches(2.5))
    return _save(presentation, tmp_path, "flow.pptx")


# ---- structure -------------------------------------------------------------------------

def test_every_slide_opens_with_the_heading_paging_cuts_at(tmp_path):
    """``office.paging`` cuts a deck at this line, so it is both the page boundary and the
    name a citation's page number resolves to. FA writes ``<!-- slide N -->``, which is a
    comment: it would name nothing and cut nothing."""
    out = reader.render(_deck(tmp_path))
    assert "## Slide 1: Findings" in out
    assert "## Slide 2: Regions" in out
    assert "## Slide 3: Growth" in out


def test_the_title_is_not_also_rendered_as_a_shape(tmp_path):
    out = reader.render(_deck(tmp_path))
    assert out.count("Findings") == 1


def test_a_slide_with_no_title_placeholder_is_still_numbered(tmp_path):
    assert "## Slide 1" in reader.render(_flow(tmp_path))


def test_bullet_levels_follow_the_paragraph_not_the_placeholder(tmp_path):
    """FA reads ``a:buChar``/``a:buAutoNum`` rather than guessing from the placeholder;
    guessing loses level-0 bullets and every numbered list."""
    out = reader.render(_deck(tmp_path))
    assert "- First point" in out
    assert "  - Sub point" in out


def test_speaker_notes_land_at_the_end_of_their_slide(tmp_path):
    """Not at the top, where FA puts them: the speaker's aside should not sit between the
    slide's title and the slide's own words."""
    out = reader.render(_deck(tmp_path))
    slide = out.split("## Slide 2")[0]
    assert slide.index("First point") < slide.index("**notes:** mention the caveat")


def test_a_chart_carries_its_real_values(tmp_path):
    """The numbers live in the embedded chart part; no text extraction reaches them."""
    out = reader.render(_deck(tmp_path))
    assert "chart COLUMN_CLUSTERED" in out
    assert "category\trevenue" in out
    assert "Q1\t1200.0" in out


def test_a_table_says_its_shape_and_separates_with_tabs(tmp_path):
    out = reader.render(_deck(tmp_path))
    assert "table 3×2" in out
    assert "region\tunits" in out
    assert "|" not in out


def test_identically_styled_shapes_collapse_into_one_group(tmp_path):
    """Five feature boxes are one object drawn five times; printing five identical style
    tags buries the five labels that actually differ."""
    out = reader.render(_flow(tmp_path))
    assert "group ×4 ROUNDED_RECTANGLE fill:#228899 18pt" in out
    assert "- `@0.5,2.0` Collect" in out


def test_connectors_are_chained_into_a_flow(tmp_path):
    """A connector's meaning is what it joins, and that is only recoverable from geometry:
    which shape's right edge sits at the arrow's left end."""
    assert "- flow: Collect → Clean → Analyse → Report" in reader.render(_flow(tmp_path))


def test_a_label_drawn_on_top_of_an_empty_box_is_one_line_not_two(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(1), Inches(1), Inches(3), Inches(2))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0xAA, 0xBB, 0xCC)
    label = slide.shapes.add_textbox(Inches(1.5), Inches(1.5), Inches(2), Inches(1))
    label.text_frame.paragraphs[0].add_run().text = "Inside"
    out = reader.render(_save(presentation, tmp_path, "pair.pptx"))
    assert out.count("Inside") == 1
    assert "fill:#AABBCC" in out                 # the box folded into the label's own line
    assert out.count("fill:#AABBCC") == 1


def test_a_shape_with_no_text_is_gathered_as_needing_a_vision_model(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8), Inches(5), Inches(1), Inches(1))
    out = reader.render(_save(presentation, tmp_path, "visual.pptx"))
    assert "▸ visual elements — needs_vlm: true" in out
    assert "OVAL vlm" in out


def test_a_slide_that_needs_no_vision_model_says_nothing_about_one(tmp_path):
    """The marker is per slide precisely so it means something. FA closes the document with
    a list of slide numbers, which here would be paged into the last slide's block and read
    as a statement about that slide."""
    out = reader.render(_deck(tmp_path))
    assert "needs_vlm" not in out.split("## Slide 2")[0]


# ---- what stays quotable ----------------------------------------------------------------

def test_two_cells_of_one_row_verify_as_one_quotation(tmp_path):
    """Under a pipe table the row folds to "|North|800|" and "North 800" is not in it."""
    out = reader.render(_deck(tmp_path))
    assert fold("North 800") in fold(out)


def test_a_line_whose_runs_differ_still_verifies_whole(tmp_path):
    """FA appends the size note to the run: ``sharply`14pt``` inside the sentence. The page
    then folds to "rose sharply`14pt` in Q3" and the sentence cannot be found in it."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(7), Inches(2)).text_frame
    paragraph = frame.paragraphs[0]
    for text, size, bold in (("Revenue rose ", 28, False),
                             ("sharply", 14, True),
                             (" in Q3.", 28, False)):
        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
    out = reader.render(_save(presentation, tmp_path, "runs.pptx"))
    assert fold("Revenue rose sharply in Q3.") in fold(out)
    assert 'bold "sharply"' in out
    assert '14pt "sharply"' in out


def test_a_line_styled_end_to_end_keeps_its_markers(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(7), Inches(1)).text_frame
    run = frame.paragraphs[0].add_run()
    run.text = "Everything here is emphasised"
    run.font.bold = True
    out = reader.render(_save(presentation, tmp_path, "allbold.pptx"))
    assert "**Everything here is emphasised**" in out
    assert fold("here is emphasised") in fold(out)


def test_a_chart_title_with_a_quote_in_it_does_not_break_its_own_tag(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["a"]
    data.add_series("s", (1.0,))
    graphic = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                                     Inches(1), Inches(1), Inches(4), Inches(3), data)
    graphic.chart.has_title = True
    graphic.chart.chart_title.text_frame.text = 'The "big"\nquestion'
    out = reader.render(_save(presentation, tmp_path, "quoted.pptx"))
    assert "title=\"The 'big' question\"" in out


# ---- edges -------------------------------------------------------------------------------

def test_the_deck_title_is_its_first_slide_title(tmp_path):
    assert reader.title(_deck(tmp_path)) == "Findings"


def test_a_deck_whose_first_slide_has_no_title_reports_none(tmp_path):
    assert reader.title(_flow(tmp_path)) == ""


def test_something_that_is_not_a_deck_is_refused_by_name(tmp_path):
    path = tmp_path / "fake.pptx"
    path.write_bytes(b"not a zip at all")
    with pytest.raises(ValueError) as caught:
        reader.render(str(path))
    assert "fake.pptx" in str(caught.value)


def test_an_empty_deck_renders_to_its_header_alone(tmp_path):
    presentation = Presentation()
    out = reader.render(_save(presentation, tmp_path, "empty.pptx"))
    assert out.startswith("<!-- pptx readout")
    assert out.strip().count("\n") == 0


def test_a_deck_pages_one_slide_at_a_time_in_the_corpus(tmp_path, monkeypatch):
    """The point of the ``## Slide N`` boundary. Cut at 3000 characters instead, a page
    would hold the tail of one slide and the top of the next, and "p7" would name nothing a
    reader could turn to."""
    monkeypatch.setenv("MISAKA_PAGEINDEX", str(tmp_path / "corpus"))
    from misaka.core.documents import index as corpus

    doc_id, pages = corpus.ingest(_deck(tmp_path))
    assert pages == 3
    located = corpus.verify_quote(doc_id, "North 800")
    assert located is not None
    assert located["page"] == 2                  # the slide the table is actually on


def test_the_deck_title_reaches_the_corpus_metadata(tmp_path, monkeypatch):
    monkeypatch.setenv("MISAKA_PAGEINDEX", str(tmp_path / "corpus"))
    from misaka.core.documents import index as corpus

    assert corpus.source_title(_deck(tmp_path)) == "Findings"
