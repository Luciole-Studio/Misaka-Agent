"""Boundary regressions from the full Office control-flow sweep."""
import builtins
import csv
import io
import zipfile
from types import SimpleNamespace
from xml.etree import ElementTree as ET

import openpyxl
import pytest
from docx import Document
from pptx import Presentation
from pptx.dml.color import RGBColor

from misaka.core.documents.office import paging, soffice
from misaka.core.documents.office import xlsx as reader
from misaka.core.tools import _office


def workbook(tmp_path):
    path = tmp_path / 'book.xlsx'
    book = openpyxl.Workbook()
    book.active.title = 'Data'
    book.active['A1'] = 'header'
    book.save(path)
    return path


@pytest.mark.parametrize('operation', ['set_cell', 'set_range'])
@pytest.mark.parametrize('text', ['=1+1', '#N/A', '00123'])
def test_explicit_text_is_not_a_formula_or_error(tmp_path, operation, text):
    path = workbook(tmp_path)
    args = ({'cell': 'A2', 'value': text, 'type': 'text'} if operation == 'set_cell'
            else {'start_cell': 'A2', 'rows': [[text]], 'types': 'text'})
    receipt = _office.run_ops(path, [{operation: args}])
    assert receipt.startswith('✓'), receipt
    cell = openpyxl.load_workbook(path).active['A2']
    assert cell.value == text
    assert cell.data_type == 's'
    assert reader._formula_lines(openpyxl.load_workbook(path).active) == ('formulas', [])


def test_cell_format_updates_preserve_other_properties(tmp_path):
    path = workbook(tmp_path)
    ops = [{'set_cell_format': {'cell_range': 'A1', 'wrap': True, 'align_v': 'top',
                              'border': {'sides': 'bottom', 'color': 'FF0000'}}},
           {'set_cell_format': {'cell_range': 'A1', 'align_h': 'right',
                              'border': {'sides': 'left', 'color': '0000FF'}}}]
    assert _office.run_ops(path, ops).startswith('✓')
    cell = openpyxl.load_workbook(path).active['A1']
    assert cell.alignment.wrap_text is True
    assert cell.alignment.vertical == 'top'
    assert cell.alignment.horizontal == 'right'
    assert cell.border.bottom.style == 'thin'
    assert cell.border.left.style == 'thin'


def test_r1c1_conversion_keeps_string_literals_and_sheet_names():
    actual = reader._r1c1('=IF(A2="A2",\'A2\'!$B$3,"B2")', 2, 2)
    assert actual == '=IF(RC[-1]="A2",\'A2\'!R3C2,"B2")'


def test_empty_string_formula_cache_is_not_uncached(tmp_path):
    path = workbook(tmp_path)
    with zipfile.ZipFile(path) as z:
        members = {n: z.read(n) for n in z.namelist()}
    namespace = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    root = ET.fromstring(members['xl/worksheets/sheet1.xml'])
    row = root.find(namespace + 'sheetData').find(namespace + 'row')
    cell = ET.SubElement(row, namespace+'c', {'r': 'B1', 't': 'str'})
    ET.SubElement(cell, namespace+'f').text = 'IF(1=1,"",1)'
    ET.SubElement(cell, namespace+'v')
    members['xl/worksheets/sheet1.xml'] = ET.tostring(root)
    with zipfile.ZipFile(path, 'w') as z:
        for n, data in members.items():
            z.writestr(n, data)
    rendered = reader.render(path)
    assert '`uncached`' not in rendered.split('-->', 1)[1]
    assert '`uncached:' not in rendered


def test_sparse_distant_cells_do_not_expand_full_worksheet(tmp_path, monkeypatch):
    path = workbook(tmp_path)
    book = openpyxl.load_workbook(path)
    book.active['XFD1048576'] = 'far corner'
    book.save(path)
    original = openpyxl.worksheet.worksheet.Worksheet.iter_rows
    def bounded(sheet, *args, **kwargs):
        assert sheet.max_row * sheet.max_column < 1_000_000, 'unbounded dense worksheet scan'
        return original(sheet, *args, **kwargs)
    monkeypatch.setattr(openpyxl.worksheet.worksheet.Worksheet, 'iter_rows', bounded)
    rendered = reader.render(path)
    assert 'far corner' in rendered and 'header' in rendered


def test_oversized_point_query_is_bounded(tmp_path, monkeypatch):
    path = workbook(tmp_path)
    original = builtins.range
    def bounded(*args):
        value = original(*args)
        assert len(value) <= reader.MAX_GRID_CELLS, 'range scanned before grid limit'
        return value
    monkeypatch.setattr(reader, 'range', bounded, raising=False)
    rendered = reader.render(path, cell_range='Data!A1:XFD1048576')
    assert 'too large' in rendered


@pytest.mark.parametrize('kind', ['table', 'pivot'])
def test_oversized_table_and_pivot_masks_are_bounded(monkeypatch, kind):
    from openpyxl.worksheet.table import Table
    sheet = openpyxl.Workbook().active
    sheet['A1'] = 'header'
    original = builtins.range
    def bounded(*args):
        value = original(*args)
        assert len(value) <= reader.MAX_GRID_CELLS, 'unbounded metadata mask'
        return value
    monkeypatch.setattr(reader, 'range', bounded, raising=False)
    if kind == 'table':
        sheet.add_table(Table(displayName='Huge', ref='A1:XFD1048576'))
        regions, mask = reader._table_regions(sheet, sheet)
    else:
        sheet._pivots = [SimpleNamespace(location=SimpleNamespace(ref='A1:XFD1048576'))]
        regions, mask = reader._pivot_regions(sheet)
    assert regions
    assert len(mask) <= len(sheet._cells)


def test_quoted_sheet_name_with_apostrophe_and_bang(tmp_path):
    path = workbook(tmp_path)
    book = openpyxl.load_workbook(path)
    book.active.title = "O'Brien! Q1"
    book.save(path)
    assert 'header' in reader.render(path, cell_range="'O''Brien! Q1'!A1")


@pytest.mark.parametrize('fmt,header', [('xlsx', '## Sheet: First'), ('pptx', '## Slide 1: First'), ('docx', '# First')])
def test_first_block_continuation_retains_title_after_readout_preamble(fmt, header):
    text = '<!-- renderer legend -->\n\n' + header + '\n\tA\tB\n1\tx\ty\n'
    blocks = paging.blocks(text, fmt)
    assert paging.resume_context(blocks[0], fmt).startswith(header + ' (continued)\n')


def test_fence_closes_only_with_same_character_and_sufficient_length():
    text = '# One\n````text\n~~~\n# Not a heading\n```\n# Still code\n````\n# Two\nbody'
    blocks = paging.blocks(text, 'docx')
    assert len(blocks) == 2
    assert '# Still code' in blocks[0]


def test_convert_does_not_return_stale_output(tmp_path, monkeypatch):
    source = tmp_path / 'source.docx'
    source.write_bytes(b'input')
    out = tmp_path / 'out'
    out.mkdir()
    old = out / 'source.pdf'
    old.write_bytes(b'old output')
    monkeypatch.setattr(soffice, 'binary', lambda: 'fixture')
    monkeypatch.setattr(soffice, '_run', lambda *args, **kwargs: True)
    meta = {}
    assert soffice.convert(source, 'pdf', into=out, meta=meta) is None
    assert old.read_bytes() == b'old output'
    assert 'no pdf produced' in meta['soffice_error']


@pytest.mark.parametrize('layout', ['merged', 'nested'])
def test_docx_replacement_visits_each_nested_paragraph_once(tmp_path, layout):
    path = tmp_path / 'document.docx'
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    if layout == 'merged':
        cell = table.cell(0, 0).merge(table.cell(0, 1))
    else:
        cell = table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0)
    cell.text = 'x'
    doc.save(path)
    receipt = _office.run_ops(path, [{'replace_text': {'find': 'x', 'replace': 'xx'}}])
    assert receipt.startswith('✓'), receipt
    from misaka.core.documents.office import docx as docx_reader
    actual = docx_reader.render(path)
    assert 'xx' in actual and 'xxxx' not in actual, actual
    assert 'replaced 1 occurrence' in receipt


def test_docx_zero_replacement_count_is_a_noop(tmp_path):
    path = tmp_path / 'document.docx'
    doc = Document(); doc.add_paragraph('original'); doc.save(path)
    receipt = _office.run_ops(path, [{'replace_text': {'find': 'original', 'replace': 'changed', 'count': 0}}])
    assert receipt.startswith('✓'), receipt
    assert Document(path).paragraphs[0].text == 'original'


def test_duplicate_slide_keeps_notes_and_background(tmp_path):
    path = tmp_path / 'deck.pptx'
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
    slide.notes_slide.notes_text_frame.text = 'Do not lose speaker notes'
    deck.save(path)
    assert _office.run_ops(path, [{'duplicate_slide': {'slide': 1}}]).startswith('✓')
    fresh = Presentation(path)
    assert fresh.slides[1].notes_slide.notes_text_frame.text == 'Do not lose speaker notes'
    assert fresh.slides[1].background.fill.fore_color.rgb == RGBColor(0x12, 0x34, 0x56)


def test_pptx_slide_zero_does_not_edit_every_slide(tmp_path):
    path = tmp_path / 'deck.pptx'
    assert _office.run_ops(path, [{'create': {'slides': [{'title': 'original'}]}}]).startswith('✓')
    before = path.read_bytes()
    receipt = _office.run_ops(path, [{'replace_text': {'slide': 0, 'find': 'original', 'replace': 'changed'}}])
    assert receipt.startswith('✗'), receipt
    assert path.read_bytes() == before


def test_csv_preserves_extra_columns_and_literal_quoting(tmp_path):
    path = tmp_path / 'ragged.csv'
    content = 'name\n"first","extra"\n"multi\nline","tail"\n'
    path.write_text(content)
    rendered = reader.render_csv(path)
    assert content.rstrip('\n') in rendered
    body = rendered.split('```\n', 1)[1]
    assert list(csv.reader(io.StringIO(body))) == list(csv.reader(io.StringIO(content)))


def test_writers_precheck_existing_office_packages(tmp_path):
    path = tmp_path / 'document.docx'
    Document().save(path)
    with zipfile.ZipFile(path) as z:
        members = {n: z.read(n) for n in z.namelist()}
    members['word/document.xml'] = b'<!DOCTYPE x [<!ENTITY x "untrusted">]>' + members['word/document.xml'].split(b'?>', 1)[1]
    with zipfile.ZipFile(path, 'w') as z:
        for n, data in members.items(): z.writestr(n, data)
    before = path.read_bytes()
    receipt = _office.run_ops(path, [{'insert_paragraph': {'text': 'new'}}])
    assert receipt.startswith('✗') and 'entities' in receipt, receipt
    assert path.read_bytes() == before


async def test_permission_checks_and_execution_share_ops_file_snapshot(tmp_path, monkeypatch):
    import json

    from misaka.core.subagent import configuration, policy
    from misaka.core.tools.office import create_office_tool_definition

    context = SimpleNamespace(workspace=str(tmp_path), permission_mode='acceptEdits')
    guard = policy.AgentPolicy(context)
    async def refresh():
        return []
    monkeypatch.setattr(policy, 'refresh_inherited_permissions', refresh)
    monkeypatch.setattr(configuration, 'permission_settings', lambda *args, **kwargs: [])
    source = tmp_path / 'ops.json'
    source.write_text(json.dumps([{'create': {'content': 'approved content'}}]))
    original = {'path': 'report.txt', 'ops': '@ops.json'}
    decision = await guard.before_tool({'toolName': 'office', 'input': original, 'toolCallId': 'audit'})
    assert decision and 'updatedInput' in decision, decision
    assert isinstance(decision['updatedInput']['ops'], list)
    source.write_text(json.dumps([{'create': {'content': 'changed after permission check'}}]))
    await create_office_tool_definition(str(tmp_path)).execute('audit', decision['updatedInput'])
    assert (tmp_path / 'report.txt').read_text() == 'approved content'


def test_cache_detects_same_size_edits_with_restored_mtime(tmp_path):
    import os

    from misaka.core.documents.office import cache
    path = tmp_path / 'document.docx'
    path.write_text('old')
    first = path.stat()
    assert cache.render_cached(path, path.read_text, workspace=tmp_path) == 'old'
    path.write_text('new')
    os.utime(path, ns=(first.st_atime_ns, first.st_mtime_ns))
    assert cache.render_cached(path, path.read_text, workspace=tmp_path) == 'new'


def test_partial_formula_cache_still_reports_resave_risk_and_errors(tmp_path, monkeypatch):
    path = workbook(tmp_path)
    monkeypatch.setattr(_office, '_recalculate', lambda *args: True)
    monkeypatch.setattr(_office._xlsx, 'formula_errors', lambda *args: ['Data!B2 #VALUE!'])
    receipt = _office.run_ops(path, [{'set_cell': {'cell': 'B2', 'value': '=1/0'}}])
    assert receipt.startswith('✓'), receipt
    assert 'uncached' in receipt
    assert 'workbook re-saved' in receipt
    assert 'Data!B2 #VALUE!' in receipt


def test_non_string_image_path_is_rejected_before_permission_checks():
    from misaka.core.tools._office.paths import resolve_ops
    with pytest.raises((TypeError, ValueError), match='path'):
        resolve_ops([{'add_image': {'image_path': 123}}], '/fixture/workspace')


def test_duplicate_chart_can_be_edited_without_changing_source_slide(tmp_path):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    path = tmp_path / 'chart.pptx'
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    data = CategoryChartData(); data.categories = ['A']; data.add_series('S', [1])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(5), Inches(4), data)
    deck.save(path)
    assert _office.run_ops(path, [{'duplicate_slide': {'slide': 1}}]).startswith('✓')
    deck = Presentation(path)
    original = deck.slides[0].shapes[0].chart
    copied = deck.slides[1].shapes[0].chart
    changed = CategoryChartData(); changed.categories = ['A']; changed.add_series('S', [9])
    copied.replace_data(changed)
    assert original.series[0].values == (1.0,)
    assert copied.series[0].values == (9.0,)
    assert original.part.chart_workbook.xlsx_part is not copied.part.chart_workbook.xlsx_part
    deck.save(path)
    reopened = Presentation(path)
    assert reopened.slides[0].shapes[0].chart.series[0].values == (1.0,)


@pytest.mark.parametrize('failure', [1, 2, 3])
@pytest.mark.parametrize('existing_exports', [False, True])
def test_every_publish_failure_restores_document_and_exports(tmp_path, monkeypatch, failure, existing_exports):
    path = tmp_path / 'document.docx'
    doc = Document(); doc.add_paragraph('before'); doc.save(path)
    exports = [tmp_path / 'one.pdf', tmp_path / 'two.pdf']
    if existing_exports:
        for export in exports: export.write_bytes(b'previous PDF')
    before = {target: target.read_bytes() if target.exists() else None for target in [path, *exports]}
    monkeypatch.setattr(soffice, 'binary', lambda: 'fixture')
    def export_pdf(source, target, **kwargs):
        from pathlib import Path
        Path(target).write_bytes(b'%PDF-1.7\nfixture')
        return True
    monkeypatch.setattr(soffice, 'export_pdf', export_pdf)
    real_replace = _office.os.replace
    count = 0
    def fault(source, target):
        nonlocal count
        if str(source).endswith(('.docx', '.pdf')):
            count += 1
            if count == failure: raise PermissionError('publish fault')
        return real_replace(source, target)
    monkeypatch.setattr(_office.os, 'replace', fault)
    receipt = _office.run_ops(path, [{'insert_paragraph': {'text': 'after'}},
                                   *[{'export_pdf': {'out': str(p)}} for p in exports]])
    assert receipt.startswith('✗') and 'publish fault' in receipt
    for target, content in before.items():
        assert (target.read_bytes() if target.exists() else None) == content
    assert not list(tmp_path.glob('.office-*'))


@pytest.mark.parametrize('platform,expected', [('darwin', 'osx'), ('linux', 'svp')])
def test_soffice_uses_platform_font_backend(tmp_path, monkeypatch, platform, expected):
    import subprocess
    seen = {}
    monkeypatch.setattr(soffice, 'binary', lambda: 'fixture')
    monkeypatch.setattr(soffice.sys, 'platform', platform)
    monkeypatch.delenv('SAL_USE_VCLPLUGIN', raising=False)
    def run(args, **kwargs):
        seen.update(kwargs['env'])
        return subprocess.CompletedProcess(args, 0, '', '')
    monkeypatch.setattr(soffice.subprocess, 'run', run)
    assert soffice._run([], timeout=1, meta={}, directory=tmp_path)
    assert seen['SAL_USE_VCLPLUGIN'] == expected
    monkeypatch.setenv('SAL_USE_VCLPLUGIN', 'custom-plugin')
    assert soffice._run([], timeout=1, meta={}, directory=tmp_path)
    assert seen['SAL_USE_VCLPLUGIN'] == 'custom-plugin'


@pytest.mark.parametrize('container', ['textbox', 'table', 'group'])
def test_pptx_replacement_crosses_runs_in_all_text_containers(tmp_path, container):
    from pptx.util import Inches
    path = tmp_path / 'text.pptx'
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    if container == 'table':
        frame = slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(4), Inches(1)).table.cell(0, 0).text_frame
    else:
        shapes = slide.shapes.add_group_shape().shapes if container == 'group' else slide.shapes
        frame = shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame
    paragraph = frame.paragraphs[0]
    paragraph.add_run().text = 'prefix '
    run = paragraph.add_run(); run.text = 'hel'; run.font.bold = True
    paragraph.add_run().text = 'lo world'
    tail = paragraph.add_run(); tail.text = ' suffix'; tail.font.italic = True
    deck.save(path)
    receipt = _office.run_ops(path, [{'replace_text': {'find': 'hello world', 'replace': 'replacement'}}])
    assert receipt.startswith('✓'), receipt
    shape = Presentation(path).slides[0].shapes[0]
    if container == 'group':
        shape = shape.shapes[0]
    frame = shape.table.cell(0, 0).text_frame if container == 'table' else shape.text_frame
    assert frame.text == 'prefix replacement suffix'
    assert frame.paragraphs[0].runs[-1].font.italic is True


def test_pptx_soft_break_is_not_a_silent_word_join(tmp_path):
    from pptx.util import Inches
    path = tmp_path / 'break.pptx'
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    paragraph = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.paragraphs[0]
    paragraph.add_run().text = 'hello'; paragraph.add_line_break(); paragraph.add_run().text = 'world'
    deck.save(path)
    receipt = _office.run_ops(path, [{'replace_text': {'find': 'helloworld', 'replace': 'wrong'}}])
    assert '0 matches' in receipt
    assert Presentation(path).slides[0].shapes[0].text == 'hello\vworld'


def test_r1c1_preserves_names_that_are_not_excel_coordinates():
    assert reader._r1c1('=ZZZ1+A0+A1', 1, 1) == '=ZZZ1+A0+RC'


def test_explicit_overwrite_can_rebuild_a_damaged_package(tmp_path):
    path = tmp_path / 'damaged.docx'
    path.write_bytes(b'not a zip')
    receipt = _office.run_ops(path, [{'create': {'blocks': [{'text': 'recovered'}]}}], overwrite=True)
    assert receipt.startswith('✓'), receipt
    assert Document(path).paragraphs[0].text == 'recovered'


async def test_permission_request_hook_rewrite_snapshots_ops(tmp_path, monkeypatch):
    import json

    from misaka.core.subagent import configuration, policy
    guard = policy.AgentPolicy(SimpleNamespace(workspace=str(tmp_path), permission_mode='default'))
    async def refresh(): return []
    monkeypatch.setattr(policy, 'refresh_inherited_permissions', refresh)
    monkeypatch.setattr(configuration, 'permission_settings', lambda *args, **kwargs: [])
    source = tmp_path / 'ops.json'
    source.write_text(json.dumps([{'create': {'content': 'hook approved'}}]))
    async def hooks(event_name, *args):
        if event_name == 'PermissionRequest':
            return [{'updated_input': {'path': 'report.txt', 'ops': '@ops.json'}, 'decision': 'allow'}]
        return []
    monkeypatch.setattr(guard, '_execute_hooks', hooks)
    decision = await guard.before_tool({'toolName': 'office', 'input': {'path': 'report.txt', 'content': 'initial'}})
    assert decision == {'updatedInput': {'path': 'report.txt', 'ops': [{'create': {'content': 'hook approved'}}]}}
    source.write_text('[]')
    assert decision['updatedInput']['ops'][0]['create']['content'] == 'hook approved'
