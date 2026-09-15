"""``read`` shows a spreadsheet as a spreadsheet, and shows the corpus' own rendering.

Before this, ``read`` decoded whatever it was handed with ``errors="replace"``: a .xlsx is
a zip, so it came back as pages of replacement characters that look like content, pass any
"is there text here" check a caller makes, and quote against nothing. The model's only
working path was to already know that ``doc_add`` exists.

The property worth protecting is not that it renders -- it is that it renders *the same
text the corpus stores*. A researcher who reads a workbook in the working directory,
copies a figure out of it, and then cites it after ``doc_add`` must have the citation
verify; two renderers would make that a coin flip.

``save_to`` from the plan is deliberately absent; see
``test_read_stays_a_read_only_tool``.
"""
from __future__ import annotations

import importlib

import pytest

from misaka.core.tools.read import create_read_tool_definition

openpyxl = importlib.import_module("openpyxl")


async def _read(cwd, **params):
    tool = create_read_tool_definition(str(cwd))
    result = await tool.execute("call-1", params, None, None, None)
    return "\n".join(part.text for part in result.content if hasattr(part, "text"))


def _book(tmp_path, name="book.xlsx"):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Revenue"
    ws["A1"], ws["B1"] = "year", "amount"
    for offset, row in enumerate(range(2, 6)):
        ws.cell(row=row, column=1, value=2020 + offset)
        ws.cell(row=row, column=2, value=(offset + 1) * 1000)
    path = tmp_path / name
    wb.save(path)
    return path


async def test_a_workbook_reads_as_a_grid_not_as_replacement_characters(tmp_path):
    _book(tmp_path)
    out = await _read(tmp_path, path="book.xlsx")
    assert "## Sheet: Revenue" in out
    assert "1\tyear\tamount" in out
    assert "�" not in out


async def test_read_and_the_corpus_agree_byte_for_byte(tmp_path, monkeypatch):
    """The whole reason the rendering is shared. A quotation copied out of ``read`` has to
    verify against the document ``doc_add`` indexed, and it can only do that if the two
    saw the same text."""
    monkeypatch.setenv("MISAKA_PAGEINDEX", str(tmp_path / "corpus"))
    from misaka.core.documents import index as corpus

    path = _book(tmp_path)
    shown = await _read(tmp_path, path="book.xlsx")
    doc_id, _pages = corpus.ingest(str(path))
    quoted = "year\tamount"
    assert quoted in shown
    assert corpus.verify_quote(doc_id, quoted) is not None


async def test_offset_pages_the_rendering_the_way_it_pages_a_text_file(tmp_path):
    _book(tmp_path)
    whole = await _read(tmp_path, path="book.xlsx")
    tail = await _read(tmp_path, path="book.xlsx", offset=4)
    assert len(tail) < len(whole)
    assert "## Sheet: Revenue" not in tail          # the head lines were skipped
    assert "2023\t4000" in tail


async def test_cell_range_reads_one_region(tmp_path):
    _book(tmp_path)
    out = await _read(tmp_path, path="book.xlsx", cell_range="Revenue!A2:B3")
    assert "2020\t1000" in out
    assert "2023\t4000" not in out


async def test_cell_range_on_something_that_is_not_a_spreadsheet_is_refused(tmp_path):
    (tmp_path / "notes.md").write_text("# notes\n", encoding="utf-8")
    with pytest.raises(RuntimeError) as caught:
        await _read(tmp_path, path="notes.md", cell_range="Sheet1!A1:B2")
    assert "cell_range" in str(caught.value)


async def test_a_csv_reads_as_a_table_with_its_columns_typed(tmp_path):
    (tmp_path / "rows.csv").write_text(
        "name,amount\n" + "".join(f"r{n},{n}\n" for n in range(1, 31)), encoding="utf-8")
    out = await _read(tmp_path, path="rows.csv")
    assert "▸ table" in out
    assert "amount: num, min=1, max=30" in out


async def test_a_word_document_reads_as_a_document(tmp_path):
    """Before this, ``read`` decoded a .docx with ``errors="replace"``: a zip came back as
    pages of replacement characters that look like content and quote against nothing."""
    _docx = importlib.import_module("docx")
    document = _docx.Document()
    document.add_heading("Method", 1)
    document.add_paragraph("We sampled two hundred records.")
    document.save(tmp_path / "paper.docx")
    out = await _read(tmp_path, path="paper.docx")
    assert "# Method" in out
    assert "We sampled two hundred records." in out
    assert "\ufffd" not in out


async def test_a_deck_reads_slide_by_slide(tmp_path):
    pptx = importlib.import_module("pptx")
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Results"
    presentation.save(tmp_path / "deck.pptx")
    out = await _read(tmp_path, path="deck.pptx")
    assert "## Slide 1: Results" in out


async def test_an_unreadable_workbook_says_what_is_wrong_with_it(tmp_path):
    (tmp_path / "fake.xlsx").write_bytes(b"not a zip at all")
    with pytest.raises(Exception) as caught:
        await _read(tmp_path, path="fake.xlsx")
    assert "not a zip archive" in str(caught.value)


async def test_a_plain_text_file_is_unaffected(tmp_path):
    (tmp_path / "notes.md").write_text("# heading\nbody line\n", encoding="utf-8")
    out = await _read(tmp_path, path="notes.md")
    assert "# heading" in out and "body line" in out


def test_read_stays_a_read_only_tool():
    """The plan asked for a ``save_to`` that writes the full rendering to a path. It is not
    here, and this is the reason: ``read`` is in ``policy.READ_ONLY_TOOLS``, which plan mode
    and the workspace guard both rely on to mean the tool cannot put bytes on disk. A read
    tool that writes would be allowed to write in plan mode. The capability is ``read`` then
    ``write``, or ``doc_add`` then ``doc_read`` for a document too large to hold.
    """
    from misaka.core.subagent.policy import READ_ONLY_TOOLS
    from misaka.core.tools.read import ReadToolInput

    assert "read" in READ_ONLY_TOOLS
    assert "save_to" not in ReadToolInput.model_fields


def test_edit_and_write_send_an_office_package_to_the_office_tool():
    """A .docx is a zip. Decoding it as text and writing the result back produces a file no
    program will open, and both tools would report success -- the model finds out only when
    someone tries to open the deliverable."""
    import asyncio
    import tempfile

    from misaka.core.tools.edit import create_edit_tool_definition
    from misaka.core.tools.write import create_write_tool_definition

    directory = tempfile.mkdtemp()

    async def run():
        for factory, tool, params in (
            (create_write_tool_definition, "write", {"path": "r.docx", "content": "x"}),
            (create_edit_tool_definition, "edit",
             {"path": "r.docx", "edits": [{"oldText": "a", "newText": "b"}]}),
        ):
            with pytest.raises(RuntimeError) as caught:
                await factory(directory).execute("c", params, None, None, None)
            assert "office" in str(caught.value)
            assert ".docx" in str(caught.value)

    asyncio.run(run())


async def test_a_csv_is_still_written_by_write(tmp_path):
    """Only the zip packages are refused: .csv and .tsv are text and always were."""
    from misaka.core.tools.write import create_write_tool_definition

    tool = create_write_tool_definition(str(tmp_path))
    await tool.execute("c", {"path": "rows.csv", "content": "a,b\n1,2\n"}, None, None, None)
    assert (tmp_path / "rows.csv").read_text(encoding="utf-8") == "a,b\n1,2\n"
