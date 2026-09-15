"""Every advertised writer operation executes against a real OOXML package.

Cases also drive the pinned upstream differential harness; no external dependencies
are installed dynamically and no user documents are used.
"""
from pathlib import Path

import pytest
from PIL import Image

from misaka.core.tools._office import docx, pptx, text, xlsx
from misaka.core.tools._office._receipt import normalise


def programs(image):
    return {
        "docx": [
            ("create", {"metadata": {"title": "Audit"}, "blocks": [
                {"type": "heading", "text": "Title", "level": 1},
                {"type": "paragraph", "text": "Anchor plain text"},
                {"type": "table", "rows": [["h1", "h2"], ["a", "b"]], "column_widths_in": [1, 2]}]}),
            ("insert_paragraph", {"after_text": "Anchor", "text": [{"text": "Added", "bold": True}], "list": {"type": "number", "level": 1}}),
            ("insert_heading", {"after_text": "Anchor", "text": "Subheading", "level": 2}),
            ("insert_table", {"after_text": "Anchor", "rows": [["A", "B"], ["1", "2"]], "cant_split": True}),
            ("replace_text", {"find": "plain", "replace": "changed"}),
            ("format_text", {"find": "Added", "bold": False, "italic": True, "size": 13, "color": "123456"}),
            ("format_paragraph", {"find": "Anchor", "align": "right", "space_after": 0, "keep_with_next": False}),
            ("add_hyperlink", {"find": "Subheading", "url": "https://example.invalid/"}),
            ("add_image", {"image_path": image, "after_text": "Anchor", "width": 0.2}),
            ("set_page_margins", {"top": 0.7, "left": 0.5}),
            ("set_page_orientation", {"orientation": "landscape"}),
            ("set_header_footer", {"header": "Running title", "footer": ""}),
            ("set_page_number", {"of_total": True, "start": 2, "fmt": "roman_lower"}),
        ],
        "xlsx": [
            ("create", {"sheets": [{"name": "S", "headers": ["label", "a", "b"], "rows": [["one", 1, 2], ["two", 3, 4]]}]}),
            ("set_cell", {"cell": "D1", "value": "5", "type": "number"}),
            ("set_range", {"start_cell": "D2", "rows": [["text", 2], ["more", 3]], "types": [["text", "number"], ["text", "number"]]}),
            ("set_cell_format", {"cell_range": "A1:C1", "bold": True, "font_name": "Arial", "font_size": 12, "fill_color": "FFFF00", "align_h": "center", "wrap": True, "border": {"style": "thin"}}),
            ("format_cells", {"cell_range": "B2", "bold": False}),
            ("add_table", {"cell_range": "A1:C3", "name": "Data"}),
            ("add_chart", {"data_range": "A1:C3", "chart_type": "bar", "title": "Values"}),
            ("clear_charts", {}),
            ("merge_cells", {"cell_range": "F1:G1"}),
            ("unmerge_cells", {"cell_range": "F1:G1"}),
            ("freeze_panes", {"cell": "B2"}),
            ("set_column_width", {"columns": {"A": 20}, "hidden": ["G"]}),
            ("set_row_height", {"rows": {"1": 30}, "hidden": [5]}),
            ("set_page_setup", {"orientation": "landscape", "fit_to_width": 1, "fit_to_height": 0, "margins": {"left": 0.2}, "paper_size": "a4", "print_area": "A1:E3", "print_title_rows": "1:1"}),
            ("add_sheet", {"sheet": "Other", "rows": [[1]]}),
            ("hide_sheet", {"sheet": "Other"}),
            ("show_sheet", {"sheet": "Other"}),
            ("set_sheet_visibility", {"sheet": "Other", "hidden": False}),
            ("rename_sheet", {"sheet": "Other", "new": "Renamed"}),
            ("delete_sheet", {"sheet": "Renamed"}),
            ("add_named_range", {"name": "Input", "cell_range": "$B$2:$C$3"}),
            ("delete_named_range", {"name": "Input"}),
            ("add_data_validation", {"cell_range": "F1:F3", "formula1": ["one", "two"], "prompt": "Pick"}),
            ("add_conditional_formatting", {"cell_range": "B2:C3", "rule_type": "cell_is", "value": 2, "fill_color": "FF0000"}),
            ("set_auto_filter", {"cell_range": "A1:C3"}),
            ("set_number_format", {"cell_range": "B2:C3", "number_format": "number2"}),
            ("add_image", {"image_path": image, "anchor_cell": "J1", "width": 10, "height": 10}),
        ],
        "pptx": [
            ("create", {"slides": [{"title": "One", "body": {"items": [{"text": "Body", "level": 0, "bullet": False}]}, "notes": "First note"}]}),
            ("add_slide", {"title": "Two", "subtitle": "Subtitle", "body": ["bullet"], "index": 1}),
            ("set_text", {"slide": 1, "placeholder": "title", "text": "Updated"}),
            ("add_textbox", {"slide": 1, "text": [{"text": "Text", "font": "Arial", "bold": True}], "x": 1, "y": 2, "w": 2, "h": 1}),
            ("add_table", {"slide": 1, "rows": [["h", "v"], ["a", "1"]], "x": 1, "y": 3, "w": 3, "h": 1}),
            ("add_image", {"slide": 1, "image_path": image, "x": 5, "y": 1, "w": 1}),
            ("set_notes", {"slide": 1, "text": "Updated note"}),
            ("replace_text", {"find": "Text", "replace": "Changed"}),
            ("add_shape", {"slide": 1, "shape": "rectangle", "text": "Box", "fill_color": "123456"}),
            ("add_chart", {"slide": 1, "categories": ["a", "b"], "series": {"S": [1, 2]}}),
            ("format_text", {"slide": 1, "find": "Changed", "italic": True, "strike": True}),
            ("duplicate_slide", {"slide": 1}),
            ("delete_slide", {"slide": 2}),
            ("set_slide_size", {"preset": "16:9"}),
        ],
        "txt": [("create", {"content": "literal"}), ("append", {"rows": ["second"]}), ("replace_text", {"find": "literal", "replace": "changed"})],
    }


@pytest.mark.parametrize("suffix,writer", [("docx", docx), ("xlsx", xlsx), ("pptx", pptx), ("txt", text)])
def test_every_operation_executes_and_package_reopens(tmp_path, suffix, writer):
    image = tmp_path / "image.png"
    Image.new("RGB", (2, 2)).save(image)
    program = programs(str(image))[suffix]
    assert set(writer.OPS) <= {op for op, _args in program}
    path = str(tmp_path / f"audit.{suffix}")
    for op, args in program:
        outcome = normalise(op, writer.write(path, op, args))
        assert outcome["ok"] and not outcome["warn"], (op, outcome)
        if suffix == "docx":
            from docx import Document
            assert Document(path).paragraphs
        elif suffix == "xlsx":
            from openpyxl import load_workbook
            book = load_workbook(path)
            assert book.sheetnames
            book.close()
        elif suffix == "pptx":
            from pptx import Presentation
            assert Presentation(path).slides
        else:
            assert Path(path).read_text()
